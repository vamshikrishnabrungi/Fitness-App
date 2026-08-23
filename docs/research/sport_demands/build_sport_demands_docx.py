#!/usr/bin/env python3
"""Build the polished DOCX edition of the Runlete sport-demands reference."""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
SOURCE = ROOT / "RUNLETE_SPORT_DEMANDS_REFERENCE.md"
OUTPUT = ROOT / "Runlete_Sport_Demands_Reference.docx"

INK = "15251F"
GREEN = "1F6A4A"
LIME = "B7F52A"
MUTED = "64736D"
BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
TABLE_FILL = "E8EEF5"
LIGHT_GREEN = "EAF5EF"
LIGHT_GRAY = "F4F6F5"
WHITE = "FFFFFF"

PAGE_WIDTH_DXA = 12240
PAGE_HEIGHT_DXA = 15840
CONTENT_WIDTH_DXA = 9360
TABLE_INDENT_DXA = 120


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color="C8D2CD", size="6") -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), size)
        tag.set(qn("w:space"), "0")
        tag.set(qn("w:color"), color)


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_cant_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    node = OxmlElement("w:cantSplit")
    tr_pr.append(node)


def set_table_fixed_width(table, widths: list[int]) -> None:
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    tbl_w = tbl_pr.find(qn("w:tblW"))
    tbl_w.set(qn("w:w"), str(CONTENT_WIDTH_DXA))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(TABLE_INDENT_DXA))
    tbl_ind.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths[idx]))
            tc_w.set(qn("w:type"), "dxa")


def configure_numbering(doc: Document) -> None:
    numbering = doc.part.numbering_part.element
    max_abstract = max((int(x.get(qn("w:abstractNumId"))) for x in numbering.findall(qn("w:abstractNum"))), default=0)
    max_num = max((int(x.get(qn("w:numId"))) for x in numbering.findall(qn("w:num"))), default=0)

    def create(abstract_id: int, num_id: int, fmt: str, text: str, font: str | None = None) -> None:
        abstract = OxmlElement("w:abstractNum")
        abstract.set(qn("w:abstractNumId"), str(abstract_id))
        multi = OxmlElement("w:multiLevelType")
        multi.set(qn("w:val"), "singleLevel")
        abstract.append(multi)
        lvl = OxmlElement("w:lvl")
        lvl.set(qn("w:ilvl"), "0")
        start = OxmlElement("w:start")
        start.set(qn("w:val"), "1")
        lvl.append(start)
        num_fmt = OxmlElement("w:numFmt")
        num_fmt.set(qn("w:val"), fmt)
        lvl.append(num_fmt)
        lvl_text = OxmlElement("w:lvlText")
        lvl_text.set(qn("w:val"), text)
        lvl.append(lvl_text)
        jc = OxmlElement("w:lvlJc")
        jc.set(qn("w:val"), "left")
        lvl.append(jc)
        ppr = OxmlElement("w:pPr")
        tabs = OxmlElement("w:tabs")
        tab = OxmlElement("w:tab")
        tab.set(qn("w:val"), "num")
        tab.set(qn("w:pos"), "540")
        tabs.append(tab)
        ppr.append(tabs)
        ind = OxmlElement("w:ind")
        ind.set(qn("w:left"), "540")
        ind.set(qn("w:hanging"), "270")
        ppr.append(ind)
        spacing = OxmlElement("w:spacing")
        spacing.set(qn("w:after"), "80")
        spacing.set(qn("w:line"), "300")
        spacing.set(qn("w:lineRule"), "auto")
        ppr.append(spacing)
        lvl.append(ppr)
        if font:
            rpr = OxmlElement("w:rPr")
            fonts = OxmlElement("w:rFonts")
            fonts.set(qn("w:ascii"), font)
            fonts.set(qn("w:hAnsi"), font)
            rpr.append(fonts)
            lvl.append(rpr)
        abstract.append(lvl)
        numbering.append(abstract)
        num = OxmlElement("w:num")
        num.set(qn("w:numId"), str(num_id))
        abstract_ref = OxmlElement("w:abstractNumId")
        abstract_ref.set(qn("w:val"), str(abstract_id))
        num.append(abstract_ref)
        numbering.append(num)

    bullet_num = max_num + 1
    number_num = max_num + 2
    create(max_abstract + 1, bullet_num, "bullet", "•", "Calibri")
    create(max_abstract + 2, number_num, "decimal", "%1.")
    doc._runlete_bullet_num = bullet_num
    doc._runlete_number_num = number_num


def apply_num(paragraph, num_id: int) -> None:
    ppr = paragraph._p.get_or_add_pPr()
    num_pr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num = OxmlElement("w:numId")
    num.set(qn("w:val"), str(num_id))
    num_pr.append(ilvl)
    num_pr.append(num)
    ppr.append(num_pr)


def add_hyperlink(paragraph, text: str, url: str, color=BLUE) -> None:
    relationship_id = paragraph.part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    rcolor = OxmlElement("w:color")
    rcolor.set(qn("w:val"), color)
    rpr.append(rcolor)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    rpr.append(underline)
    rfonts = OxmlElement("w:rFonts")
    rfonts.set(qn("w:ascii"), "Calibri")
    rfonts.set(qn("w:hAnsi"), "Calibri")
    rpr.append(rfonts)
    run.append(rpr)
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.append(text_node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


INLINE_RE = re.compile(r"(\*\*.+?\*\*|`.+?`|\[[^\]]+\]\([^)]+\))")


def add_inline(paragraph, text: str, *, base_bold=False, base_color=None, base_size=11) -> None:
    pos = 0
    for match in INLINE_RE.finditer(text):
        if match.start() > pos:
            run = paragraph.add_run(text[pos : match.start()])
            run.bold = base_bold
            run.font.size = Pt(base_size)
            if base_color:
                run.font.color.rgb = RGBColor.from_string(base_color)
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            run.bold = True
            run.font.size = Pt(base_size)
            if base_color:
                run.font.color.rgb = RGBColor.from_string(base_color)
        elif token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            run.font.name = "Consolas"
            run._element.rPr.rFonts.set(qn("w:ascii"), "Consolas")
            run._element.rPr.rFonts.set(qn("w:hAnsi"), "Consolas")
            run.font.size = Pt(base_size - 0.5)
            run.font.color.rgb = RGBColor.from_string(DARK_BLUE)
        else:
            label, url = re.match(r"\[([^\]]+)\]\(([^)]+)\)", token).groups()
            add_hyperlink(paragraph, label, url)
        pos = match.end()
    if pos < len(text):
        run = paragraph.add_run(text[pos:])
        run.bold = base_bold
        run.font.size = Pt(base_size)
        if base_color:
            run.font.color.rgb = RGBColor.from_string(base_color)


def set_run_font(run, name="Calibri", size=11, color=INK, bold=None, italic=None) -> None:
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def configure_styles(doc: Document) -> None:
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    values = {
        "Heading 1": (16, BLUE, 18, 10),
        "Heading 2": (13, BLUE, 14, 7),
        "Heading 3": (12, DARK_BLUE, 10, 5),
    }
    for name, (size, color, before, after) in values.items():
        style = styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.keep_together = True


def set_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Page ")
    set_run_font(run, size=9, color=MUTED)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    paragraph._p.append(fld)


def configure_sections(doc: Document) -> None:
    for section in doc.sections:
        section.page_width = Inches(8.5)
        section.page_height = Inches(11)
        section.top_margin = Inches(1)
        section.right_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.header_distance = Inches(0.492)
        section.footer_distance = Inches(0.492)

        header = section.header
        p = header.paragraphs[0]
        for run in list(p.runs):
            p._p.remove(run._r)
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_after = Pt(3)
        r = p.add_run("RUNLETE  /  SPORT DEMANDS REFERENCE")
        set_run_font(r, size=8.5, color=MUTED, bold=True)
        ppr = p._p.get_or_add_pPr()
        pbdr = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        bottom.set(qn("w:val"), "single")
        bottom.set(qn("w:sz"), "8")
        bottom.set(qn("w:space"), "4")
        bottom.set(qn("w:color"), "B8C7C0")
        pbdr.append(bottom)
        ppr.append(pbdr)

        footer = section.footer
        fp = footer.paragraphs[0]
        for run in list(fp.runs):
            fp._p.remove(run._r)
        for child in list(fp._p):
            if child.tag == qn("w:fldSimple"):
                fp._p.remove(child)
        set_page_number(fp)


def add_cover(doc: Document) -> None:
    for _ in range(7):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(12)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(18)
    r = p.add_run("RUNLETE RESEARCH")
    set_run_font(r, size=11, color=GREEN, bold=True)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(10)
    r = p.add_run("Sport Demands Reference")
    set_run_font(r, size=30, color=INK, bold=True)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run("What each launch sport requires before Runlete selects exercises or builds programs")
    set_run_font(r, size=14, color=GREEN)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(28)
    r = p.add_run("12 sports  •  event and position aware  •  evidence-informed")
    set_run_font(r, size=10.5, color=MUTED, bold=True)

    rule = doc.add_paragraph()
    rule.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = rule.add_run("━")
    set_run_font(r, size=22, color=LIME, bold=True)

    for _ in range(4):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(12)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run("Version 1.0  |  13 August 2026")
    set_run_font(r, size=11, color=INK, bold=True)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Research reference — not an exercise library, workout template or clinical guideline")
    set_run_font(r, size=9.5, color=MUTED, italic=True)
    doc.add_page_break()


def add_contents(doc: Document) -> None:
    p = doc.add_paragraph(style="Heading 1")
    p.add_run("Contents")
    entries = [
        "Purpose and evidence model",
        "Universal demand dimensions and metric model",
        "Badminton",
        "Basketball",
        "Boxing",
        "Cricket",
        "Cycling",
        "Football",
        "Mixed Martial Arts",
        "Running",
        "Swimming",
        "Tennis",
        "Volleyball",
        "HYROX",
        "Cross-sport comparison and database implications",
        "Monitoring framework, research gaps and sources",
    ]
    table = doc.add_table(rows=len(entries), cols=2)
    widths = [1000, 8360]
    set_table_fixed_width(table, widths)
    set_table_borders(table, color=WHITE, size="0")
    for idx, entry in enumerate(entries, start=1):
        set_cant_split(table.rows[idx - 1])
        left, right = table.rows[idx - 1].cells
        left.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
        right.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
        set_cell_margins(left, 40, 40, 40, 80)
        set_cell_margins(right, 40, 80, 40, 40)
        p1 = left.paragraphs[0]
        p1.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        r = p1.add_run(f"{idx:02d}")
        set_run_font(r, size=9.5, color=GREEN, bold=True)
        p2 = right.paragraphs[0]
        r = p2.add_run(entry)
        set_run_font(r, size=10.5, color=INK, bold=idx in (1, 2, 15, 16))
    doc.add_page_break()


def parse_table(lines: list[str], start: int) -> tuple[list[list[str]], int]:
    rows = []
    idx = start
    while idx < len(lines) and lines[idx].strip().startswith("|"):
        raw = lines[idx].strip().strip("|")
        cells = [c.strip() for c in raw.split("|")]
        rows.append(cells)
        idx += 1
    if len(rows) >= 2 and all(re.fullmatch(r":?-{3,}:?", cell) for cell in rows[1]):
        rows.pop(1)
    return rows, idx


def choose_widths(rows: list[list[str]]) -> list[int]:
    cols = len(rows[0])
    if cols == 2:
        first = max(len(r[0]) for r in rows)
        return [2700, 6660] if first > 28 else [2200, 7160]
    if cols == 3:
        return [1000, 2700, 5660]
    if cols == 4:
        return [1500, 2200, 2600, 3060]
    base = CONTENT_WIDTH_DXA // cols
    widths = [base] * cols
    widths[-1] += CONTENT_WIDTH_DXA - sum(widths)
    return widths


def add_table(doc: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    cols = len(rows[0])
    table = doc.add_table(rows=len(rows), cols=cols)
    widths = choose_widths(rows)
    set_table_fixed_width(table, widths)
    set_table_borders(table)
    set_repeat_table_header(table.rows[0])
    for r_idx, source_row in enumerate(rows):
        row = table.rows[r_idx]
        set_cant_split(row)
        row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
        for c_idx, value in enumerate(source_row):
            cell = row.cells[c_idx]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            set_cell_margins(cell)
            if r_idx == 0:
                set_cell_shading(cell, TABLE_FILL)
            elif r_idx % 2 == 0:
                set_cell_shading(cell, LIGHT_GRAY)
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.08
            # Keep compact priority tables together. Longer reference tables may
            # span pages with a repeated header so they do not leave a near-empty page.
            if len(rows) <= 10 and r_idx < len(rows) - 1:
                p.paragraph_format.keep_with_next = True
            add_inline(p, value, base_bold=r_idx == 0, base_color=INK, base_size=9.2 if cols >= 3 else 9.6)
    after = doc.add_paragraph()
    after.paragraph_format.space_before = Pt(4)
    after.paragraph_format.space_after = Pt(4)


def add_code_block(doc: Document, content: list[str]) -> None:
    table = doc.add_table(rows=1, cols=1)
    set_table_fixed_width(table, [CONTENT_WIDTH_DXA])
    set_table_borders(table, color="CBD8D2")
    cell = table.cell(0, 0)
    set_cell_shading(cell, LIGHT_GREEN)
    set_cell_margins(cell, 140, 180, 140, 180)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.05
    r = p.add_run("\n".join(content))
    set_run_font(r, name="Consolas", size=9.2, color=INK)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(4)


def build_document() -> None:
    doc = Document()
    configure_styles(doc)
    configure_numbering(doc)
    configure_sections(doc)
    doc.core_properties.title = "Runlete Sport Demands Reference"
    doc.core_properties.subject = "Evidence-informed physical requirements and metrics for Runlete launch sports"
    doc.core_properties.author = "Runlete"
    doc.core_properties.keywords = "sport demands, physical preparation, metrics, evidence, Runlete"

    add_cover(doc)
    add_contents(doc)

    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    idx = 0
    in_code = False
    code_lines: list[str] = []
    first_title_skipped = False
    compact_sources = False
    while idx < len(lines):
        line = lines[idx].rstrip()
        stripped = line.strip()

        if stripped.startswith("```"):
            if not in_code:
                in_code = True
                code_lines = []
            else:
                add_code_block(doc, code_lines)
                in_code = False
            idx += 1
            continue
        if in_code:
            code_lines.append(line)
            idx += 1
            continue
        if not stripped or stripped == "---":
            idx += 1
            continue
        if stripped.startswith("|"):
            rows, idx = parse_table(lines, idx)
            add_table(doc, rows)
            continue
        if stripped.startswith("# "):
            if not first_title_skipped:
                first_title_skipped = True
                idx += 1
                continue
            p = doc.add_paragraph(style="Heading 1")
            add_inline(p, stripped[2:], base_bold=True, base_size=16)
            if stripped[2:].strip() == "Selected source register":
                compact_sources = True
            idx += 1
            continue
        if stripped.startswith("## "):
            p = doc.add_paragraph(style="Heading 2")
            add_inline(p, stripped[3:], base_bold=True, base_size=13)
            idx += 1
            continue
        if stripped.startswith("### "):
            p = doc.add_paragraph(style="Heading 3")
            add_inline(p, stripped[4:], base_bold=True, base_size=12)
            idx += 1
            continue
        if stripped.startswith("- "):
            p = doc.add_paragraph()
            apply_num(p, doc._runlete_bullet_num)
            p.paragraph_format.space_after = Pt(0 if compact_sources else 4)
            p.paragraph_format.line_spacing = 1.0 if compact_sources else 1.25
            add_inline(p, stripped[2:], base_size=8.2 if compact_sources else 11)
            idx += 1
            continue
        if re.match(r"^\d+\.\s", stripped):
            p = doc.add_paragraph()
            apply_num(p, doc._runlete_number_num)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.line_spacing = 1.25
            add_inline(p, re.sub(r"^\d+\.\s+", "", stripped))
            idx += 1
            continue

        # Merge ordinary wrapped lines into one paragraph.
        paragraph_lines = [stripped]
        idx += 1
        while idx < len(lines):
            nxt = lines[idx].strip()
            if not nxt:
                break
            if nxt.startswith(("#", "- ", "|", "```")) or re.match(r"^\d+\.\s", nxt) or nxt == "---":
                break
            paragraph_lines.append(nxt)
            idx += 1
        p = doc.add_paragraph()
        add_inline(p, " ".join(paragraph_lines))

    configure_sections(doc)
    doc.save(OUTPUT)


if __name__ == "__main__":
    build_document()
    print(OUTPUT)
