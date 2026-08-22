"""Runlete Task-1 model audit.

This module performs a read-only audit of the repository's sport knowledge,
exercise audit outputs, planning fixtures, generator contracts, and the saved
exercise-review workbook.  It never writes to MongoDB or PostgreSQL.  The only
writes are deterministic reports under ``Research materials/audits/model``.

The classifications produced here are migration *candidates*.  They are not
scientific, coaching, physiotherapy, medical, or generator approval.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import re
from collections import Counter, defaultdict
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Dict, Iterable, Iterator, List, Mapping, MutableMapping, Optional, Sequence, Tuple


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_OUTPUT = ROOT / "Research materials" / "audits" / "model"
DEFAULT_WORKBOOK = Path("/Users/vamshikrishna/Downloads/Exercies-data-Evidence-Reviewed.xlsx")

AUDIT_VERSION = "runlete_model_audit_v1"
ALLOWED_DISPOSITIONS = {"retain", "rewrite", "merge", "quarantine", "reject"}

LAUNCH_SPORTS = {
    "badminton",
    "basketball",
    "boxing",
    "cricket",
    "cycling",
    "soccer",
    "mma",
    "running_endurance",
    "swimming",
    "tennis",
    "volleyball",
}
DEFERRED_SPORTS = {"hyrox"}
EXCLUDED_RESEARCH_SPORTS = {"kickboxing", "wrestling"}

SPORT_ALIASES = {
    "football": "soccer",
    "soccer_football": "soccer",
    "running": "running_endurance",
    "runner": "running_endurance",
    "endurance_running": "running_endurance",
    "combat_sports": "mma",
    "mixed_martial_arts": "mma",
}

COLLECTION_KEYS = {
    "sport_profiles",
    "sport_roles",
    "sport_training_rules",
    "planning_rules",
    "sport_teaching_progressions",
    "sport_skill_assessments",
    "sport_level_transition_rules",
    "running_workouts",
    "running_plan_rules",
    "macro_plan_templates",
    "competition_week_rules",
    "training_protocols",
    "exercise_library",
    "primary_exercise_library",
    "exercise_variation_library",
    "exercise_progression_graph",
    "mobility_drills",
    "source_sections",
}

RUNTIME_COLLECTIONS = {
    "sport_profiles",
    "sport_roles",
    "sport_training_rules",
    "planning_rules",
    "sport_teaching_progressions",
    "sport_skill_assessments",
    "sport_level_transition_rules",
    "macro_plan_templates",
    "competition_week_rules",
    "training_protocols",
    "exercise_library",
    "primary_exercise_library",
    "exercise_variation_library",
    "exercise_progression_graph",
}

PHYSICAL_TERMS = {
    "strength", "power", "force", "sprint", "acceleration", "velocity",
    "deceleration", "change of direction", "agility", "plyometric", "jump",
    "landing", "aerobic", "anaerobic", "conditioning", "endurance", "speed",
    "rotation", "anti-rotation", "isometric", "eccentric", "concentric",
    "muscle", "tendon", "load", "fatigue", "recovery", "work:rest",
    "work to rest", "rpe", "rir", "sets", "reps", "interval", "running",
    "calf", "soleus", "hamstring", "adductor", "shoulder", "trunk",
}
TECHNICAL_TACTICAL_TERMS = {
    "tactic", "game iq", "fight iq", "stroke mechanics", "shooting",
    "passing", "dribbling", "ball mastery", "batting technique", "bowling technique",
    "serve technique", "forehand", "backhand", "ringcraft", "guard", "jab",
    "cross mechanics", "rules", "scoring", "formation", "offense", "defense",
}
MEDICAL_REHAB_TERMS = {
    "rehabilitation", "rehab", "diagnosis", "return to play", "return-to-play",
    "clinical protocol", "post-operative", "postoperative", "treatment",
    "injury protocol", "medical clearance", "pathology",
}

SOURCE_TIER_HINTS: Sequence[Tuple[int, Sequence[str]]] = (
    (1, ("official_rule", "official_federation", "official_medical", "public_health", "position_stand", "consensus")),
    (2, ("systematic_review", "meta_analysis", "overview_of_reviews", "position statement", "position_stand")),
    (3, ("randomized", "controlled_trial", "biomechanical", "time_motion", "epidemiology", "cohort", "primary_study")),
    (4, ("peer_reviewed_review", "narrative_review", "clinical_commentary", "sports_science_review")),
    (5, ("coaching", "education", "case_study", "practitioner")),
)

DIMENSIONS: Mapping[str, Sequence[str]] = {
    "force_requirements": ("force", "impulse", "rate of force", "rfd", "power", "strength"),
    "force_direction": ("horizontal", "vertical", "lateral", "rotational", "multiplanar", "direction of force"),
    "contraction_type": ("concentric", "eccentric", "isometric", "stretch-shortening", "stretch shortening"),
    "movement_speed": ("velocity", "movement speed", "contact time", "sprint speed", "cadence"),
    "effort_duration": ("duration", "seconds", "minutes", "bout", "round", "rally", "spell"),
    "recovery_between_efforts": ("recovery", "rest", "work:rest", "work to rest", "inter-repetition"),
    "joint_positions": ("joint angle", "range of motion", "hip", "knee", "ankle", "shoulder", "elbow", "wrist"),
    "repeated_effort_demands": ("repeated sprint", "repeat effort", "repeat high-intensity", "work capacity", "late-round"),
    "competition_schedule": ("competition week", "match week", "fixture", "tournament", "race week", "fight camp", "taper"),
    "commonly_loaded_tissues": ("hamstring", "adductor", "calf", "achilles", "shoulder", "wrist", "lumbar", "tendon"),
    "position_event_demands": ("position", "role", "event", "distance", "stroke", "discipline", "batter", "bowler"),
    "energy_systems": ("aerobic", "anaerobic", "alactic", "glycolytic", "vo2", "threshold"),
    "external_load_integration": ("external load", "practice load", "sport practice", "match load", "bowling load", "serve volume"),
    "monitoring_assessment": ("assessment", "monitor", "test", "rpe", "readiness", "benchmark", "pain response"),
    "progression_deload": ("progression", "progress", "deload", "regress", "volume reduction"),
}

REQUIRED_OUTPUTS = (
    "audit_manifest.json",
    "audit_summary.md",
    "source_inventory.jsonl",
    "record_dispositions.jsonl",
    "sport_coverage_matrix.json",
    "evidence_provenance.json",
    "duplicates_contradictions.json",
    "current_generator_data_flow.md",
    "input_output_relevance.md",
    "running_gap_report.md",
    "cross_sport_architecture_requirements.md",
    "migration_eligibility.md",
    "implementation_backlog.md",
)


def _norm(value: Any) -> str:
    return re.sub(r"[^a-z0-9]+", "_", str(value or "").lower()).strip("_")


def normalize_sport(value: Any) -> Optional[str]:
    token = _norm(value)
    if not token:
        return None
    return SPORT_ALIASES.get(token, token)


def _stable_hash(value: Any) -> str:
    payload = json.dumps(value, sort_keys=True, ensure_ascii=False, default=str, separators=(",", ":"))
    return hashlib.sha256(payload.encode("utf-8")).hexdigest()


def _record_text(record: Mapping[str, Any]) -> str:
    return json.dumps(record, sort_keys=True, ensure_ascii=False, default=str).lower()


PROVENANCE_KEYS = {
    "source_refs", "sources", "evidence_sources", "knowledge_sources", "source_registry",
    "metadata", "source_pack_id", "source_book_id", "source_book_ids", "draft_file",
    "ingestion_method", "created_at", "updated_at", "version", "url",
}


def _substantive_text(record: Mapping[str, Any]) -> str:
    def strip(value: Any) -> Any:
        if isinstance(value, dict):
            return {key: strip(item) for key, item in value.items() if key not in PROVENANCE_KEYS}
        if isinstance(value, list):
            return [strip(item) for item in value]
        return value

    return json.dumps(strip(dict(record)), sort_keys=True, ensure_ascii=False, default=str).lower()


def _listify(value: Any) -> List[Any]:
    if value in (None, "", [], {}):
        return []
    return value if isinstance(value, list) else [value]


def _existing_id(record: Mapping[str, Any]) -> Optional[str]:
    for key in ("id", "source_id", "review_id", "asset_id", "rule_id", "template_id"):
        value = record.get(key)
        if value not in (None, ""):
            return str(value)
    return None


def _sport_from_path(path: Path) -> Optional[str]:
    parts = [normalize_sport(part) for part in path.parts]
    for part in parts:
        if part in LAUNCH_SPORTS | DEFERRED_SPORTS | EXCLUDED_RESEARCH_SPORTS | {"soccer"}:
            return part
    if "football" in path.parts:
        return "soccer"
    if "running" in path.parts:
        return "running_endurance"
    return None


def _record_sport(record: Mapping[str, Any], path: Path, parent_sport: Optional[str]) -> Optional[str]:
    for key in ("sport", "primary_sport"):
        sport = normalize_sport(record.get(key))
        if sport:
            return sport
    applies = _listify(record.get("applies_to"))
    for value in applies:
        sport = normalize_sport(value)
        if sport in LAUNCH_SPORTS | DEFERRED_SPORTS | EXCLUDED_RESEARCH_SPORTS:
            return sport
    return parent_sport or _sport_from_path(path)


def _record_domain(record: Mapping[str, Any]) -> str:
    for key in ("domain", "category", "rule_type", "workout_type", "exercise_family", "section"):
        value = record.get(key)
        if value not in (None, ""):
            return str(value)
    return "unspecified"


def _source_refs(record: Mapping[str, Any]) -> List[Dict[str, Any]]:
    refs: List[Dict[str, Any]] = []
    for key in ("source_refs", "sources", "evidence_sources", "knowledge_sources"):
        for value in _listify(record.get(key)):
            if isinstance(value, dict):
                refs.append(dict(value))
            elif value:
                refs.append({"title": str(value), "reference_only": True})
    return refs


def _content_scope(record_type: str, domain: str, text: str) -> str:
    if any(term in text for term in MEDICAL_REHAB_TERMS):
        return "medical_or_rehabilitation"
    technical_hits = sum(term in text for term in TECHNICAL_TACTICAL_TERMS)
    physical_hits = sum(term in text for term in PHYSICAL_TERMS)
    if record_type in {"sport_teaching_progressions", "sport_skill_assessments"} and technical_hits >= physical_hits:
        return "technical_or_tactical"
    if technical_hits >= 2 and physical_hits == 0:
        return "technical_or_tactical"
    if physical_hits:
        return "physical_preparation"
    if record_type in {"source_sections", "evidence_source"}:
        return "evidence_or_provenance"
    return "unclassified"


def _layer(record_type: str, scope: str) -> str:
    if record_type in {
        "exercise_library", "primary_exercise_library", "exercise_variation_library",
        "mobility_drills", "exercise_audit_record", "workbook_exercise_candidate",
        "sport_teaching_progressions",
    }:
        return "layer_1_method_library"
    if record_type in {"exercise_progression_graph", "exercise_effect", "exercise_relation"}:
        return "layer_2_exercise_effects"
    if record_type in {"sport_profiles", "sport_roles", "sport_skill_assessments"}:
        return "layer_3_sport_demands"
    if record_type in {
        "sport_training_rules", "planning_rules", "running_plan_rules",
        "sport_level_transition_rules", "training_protocols", "competition_week_rules",
    }:
        return "layer_4_prescriptions_constraints"
    if record_type in {"macro_plan_templates", "running_workouts", "program_fixture", "workout_fixture"}:
        return "layer_5_program_recipes"
    if scope == "evidence_or_provenance" or record_type in {"source_sections", "evidence_source"}:
        return "governance_evidence"
    return "unmapped"


def _bounded_prescription(text: str) -> bool:
    has_number = bool(re.search(r"\b\d+(?:\.\d+)?\b", text))
    has_dose = any(term in text for term in ("sets", "reps", "seconds", "minutes", "rpe", "rir", "%", "week", "rest", "duration"))
    return has_number and has_dose


def _claim_support_class(record: Mapping[str, Any], text: str) -> str:
    """Preliminary provenance class; never a scientific approval decision."""
    refs = _source_refs(record)
    quantified = bool(re.search(
        r"\b\d+(?:\.\d+)?\s*(?:%|s|sec|seconds|min|minutes|m|km|kg|n|w|w/kg|week|weeks|reps|sets)\b",
        text,
    ))
    measured_terms = ("measured", "observed", "time-motion", "biomechan", "force", "velocity", "duration", "incidence")
    if refs and quantified and any(term in text for term in measured_terms):
        return "directly_measured_candidate"
    if refs:
        return "evidence_linked_implication_candidate"
    if record.get("source_pack_id") or record.get("source_book_id") or record.get("source_book_ids"):
        return "source_pack_inference_candidate"
    return "unsupported_or_unlinked_assertion"


def _initial_disposition(
    *,
    sport: Optional[str],
    record_type: str,
    scope: str,
    path: Path,
    has_id: bool,
    evidence_linked: bool,
    text: str,
) -> Tuple[str, str]:
    path_text = path.as_posix().lower()
    if "/test_outputs/" in path_text:
        return "quarantine", "generated or captured test artifact; not canonical knowledge"
    if sport in DEFERRED_SPORTS:
        return "quarantine", "Hyrox is deferred until a dedicated evidence package passes audit"
    if sport in EXCLUDED_RESEARCH_SPORTS:
        return "quarantine", "research exists but sport is outside the locked launch set"
    if sport and sport not in LAUNCH_SPORTS:
        return "quarantine", "record belongs to a sport outside the locked launch set"
    if scope == "technical_or_tactical":
        return "reject", "technical/tactical content is outside physical-preparation generation scope"
    if scope == "medical_or_rehabilitation":
        return "quarantine", "medical or rehabilitation content requires a separately governed clinical pathway"
    if "draft" in path_text:
        return "rewrite", "research draft must be normalized and reviewed before runtime use"
    if record_type in {"source_sections", "evidence_source"}:
        return "retain", "retain as evidence/provenance candidate; not generator content"
    if not has_id:
        return "rewrite", "record lacks a stable source identifier"
    if scope == "physical_preparation" and evidence_linked:
        return "retain", "physical-preparation candidate with traceable evidence; human review pending"
    if scope == "physical_preparation":
        return "rewrite", "physical-preparation candidate lacks direct evidence linkage or executable structure"
    if record_type in RUNTIME_COLLECTIONS:
        return "rewrite", "runtime-consumed record is not clearly scoped to physical preparation"
    return "quarantine", "unclassified candidate requires manual product-scope review"


@dataclass(frozen=True)
class RawRecord:
    path: Path
    locator: str
    record_type: str
    record: Dict[str, Any]
    parent_sport: Optional[str]
    source_kind: str


def _iter_collection_records(path: Path, data: Any) -> Iterator[RawRecord]:
    path_sport = _sport_from_path(path)
    if "/test_outputs/" in path.as_posix().lower():
        record = data if isinstance(data, dict) else {"value": data}
        yield RawRecord(path, "$", "test_output_artifact", record, path_sport, "test_output")
        return

    emitted: set[Tuple[str, int]] = set()

    def emit_list(record_type: str, values: Any, prefix: str, source_kind: str) -> Iterator[RawRecord]:
        if not isinstance(values, list):
            return
        for index, value in enumerate(values):
            if not isinstance(value, dict):
                value = {"value": value}
            marker = (prefix, index)
            if marker in emitted:
                continue
            emitted.add(marker)
            yield RawRecord(path, f"{prefix}[{index}]", record_type, dict(value), path_sport, source_kind)

    if isinstance(data, dict):
        collections = data.get("collections")
        if isinstance(collections, dict):
            for key, values in collections.items():
                yield from emit_list(str(key), values, f"$.collections.{key}", "compiled_collection")

        for key, values in data.items():
            if key == "collections":
                continue
            if key in COLLECTION_KEYS and isinstance(values, list):
                yield from emit_list(key, values, f"$.{key}", "compiled_collection")

        for key in ("source_registry", "knowledge_sources"):
            values = data.get(key)
            if isinstance(values, list):
                yield from emit_list("evidence_source", values, f"$.{key}", "evidence_registry")

        is_draft = "draft" in path.name.lower() or "/drafts/" in path.as_posix().lower()
        if is_draft:
            ignored = {"metadata", "source_refs", "backend_records_to_create"}
            for key, values in data.items():
                if key in ignored or key in COLLECTION_KEYS or key == "collections":
                    continue
                if isinstance(values, list) and values and all(isinstance(item, (dict, str)) for item in values):
                    yield from emit_list(f"draft_{key}", values, f"$.{key}", "research_draft")

    elif isinstance(data, list):
        yield from emit_list("untyped_list_record", data, "$", "untyped_json")


def discover_json_records() -> Tuple[List[RawRecord], List[Dict[str, str]], Dict[str, int]]:
    records: List[RawRecord] = []
    errors: List[Dict[str, str]] = []
    counts = {"files_discovered": 0, "files_parsed": 0}
    roots = (
        ROOT / "Research materials" / "legacy" / "sport research",
        ROOT / "Research materials" / "legacy" / "ai generation experiments",
    )
    for root in roots:
        for path in sorted(root.rglob("*.json")):
            counts["files_discovered"] += 1
            try:
                data = json.loads(path.read_text(encoding="utf-8"))
            except Exception as exc:  # pragma: no cover - exercised by real data
                error = {"path": path.relative_to(ROOT).as_posix(), "error": str(exc)}
                errors.append(error)
                raw_bytes = path.read_bytes()
                records.append(RawRecord(
                    path=path,
                    locator="$",
                    record_type="invalid_json_artifact",
                    record={
                        "parse_error": str(exc),
                        "file_size_bytes": len(raw_bytes),
                        "file_sha256": hashlib.sha256(raw_bytes).hexdigest(),
                    },
                    parent_sport=_sport_from_path(path),
                    source_kind="invalid_json_artifact",
                ))
                continue
            counts["files_parsed"] += 1
            records.extend(_iter_collection_records(path, data))
    return records, errors, counts


def _inventory_row(raw: RawRecord) -> Dict[str, Any]:
    record = raw.record
    source_path = raw.path.relative_to(ROOT).as_posix()
    existing_id = _existing_id(record)
    audit_id = existing_id or f"audit_{_stable_hash([source_path, raw.locator])[:20]}"
    sport = _record_sport(record, raw.path, raw.parent_sport)
    text = _substantive_text(record)
    domain = _record_domain(record)
    refs = _source_refs(record)
    evidence_linked = bool(refs or record.get("source_pack_id") or record.get("source_book_id") or record.get("source_book_ids"))
    scope = _content_scope(raw.record_type, domain, text)
    layer = _layer(raw.record_type, scope)
    disposition, reason = _initial_disposition(
        sport=sport,
        record_type=raw.record_type,
        scope=scope,
        path=raw.path,
        has_id=bool(existing_id),
        evidence_linked=evidence_linked,
        text=text,
    )
    semantic_payload = {
        key: record.get(key)
        for key in (
            "name", "title", "sport", "domain", "category", "condition", "rule",
            "purpose", "progression_rule", "recommended_action", "blocked_action",
            "learning_goal", "technical_focus", "structure",
        )
        if record.get(key) not in (None, "", [], {})
    }
    return {
        "audit_record_id": f"ari_{_stable_hash([source_path, raw.locator, audit_id])[:24]}",
        "source_id": existing_id,
        "audit_only_id": audit_id if not existing_id else None,
        "id_origin": "source" if existing_id else "synthetic_audit_only",
        "source_path": source_path,
        "source_locator": raw.locator,
        "source_kind": raw.source_kind,
        "record_type": raw.record_type,
        "sport": sport,
        "in_launch_scope": sport in LAUNCH_SPORTS or sport is None,
        "domain": domain,
        "content_scope": scope,
        "knowledge_layer": layer,
        "runtime_collection_candidate": raw.record_type in RUNTIME_COLLECTIONS,
        "evidence_linked": evidence_linked,
        "evidence_reference_count": len(refs),
        "has_stable_source_id": bool(existing_id),
        "has_bounded_prescription": _bounded_prescription(text),
        "claim_support_class": _claim_support_class(record, text),
        "claim_support_classification_status": "automated_preliminary_pending_manual_verification",
        "payload_hash": _stable_hash(record),
        "semantic_hash": _stable_hash(semantic_payload) if semantic_payload else None,
        "disposition": disposition,
        "disposition_reason": reason,
        "audit_confidence": "automated_preliminary",
        "scientific_review_status": "pending_claim_level_review",
        "human_review_required": True,
        "generator_eligible": False,
        "migration_eligible": False,
        "record": record,
    }


def load_exercise_audit_records() -> Tuple[List[Dict[str, Any]], Dict[str, Any]]:
    path = ROOT / "Research materials" / "audits" / "exercise library" / "exercise_inventory.json"
    if not path.exists():
        return [], {"present": False, "path": path.relative_to(ROOT).as_posix()}
    data = json.loads(path.read_text(encoding="utf-8"))
    rows: List[Dict[str, Any]] = []
    mapping = {
        "core": "retain",
        "contextual": "retain",
        "rewrite_required": "rewrite",
        "specialist": "quarantine",
        "quarantine": "quarantine",
        "reject": "reject",
        "merge_duplicate": "merge",
    }
    for index, record in enumerate(data.get("records") or []):
        source_id = str(record.get("source_id") or record.get("id") or f"exercise_audit_{index}")
        disposition = mapping.get(str(record.get("disposition") or ""), "rewrite")
        rows.append({
            "audit_record_id": f"ari_{_stable_hash([path.as_posix(), index, source_id])[:24]}",
            "source_id": source_id,
            "audit_only_id": None,
            "id_origin": "source",
            "source_path": path.relative_to(ROOT).as_posix(),
            "source_locator": f"$.records[{index}]",
            "source_kind": "exercise_audit_output",
            "record_type": "exercise_audit_record",
            "sport": None,
            "in_launch_scope": True,
            "domain": str(record.get("record_type") or record.get("movement_pattern") or "exercise"),
            "content_scope": "physical_preparation",
            "knowledge_layer": "layer_1_method_library",
            "runtime_collection_candidate": False,
            "evidence_linked": bool(record.get("source_refs") or record.get("source_collection")),
            "evidence_reference_count": len(_listify(record.get("source_refs"))),
            "has_stable_source_id": True,
            "has_bounded_prescription": False,
            "claim_support_class": "source_pack_inference_candidate",
            "claim_support_classification_status": "carried_from_prior_audit_pending_manual_verification",
            "payload_hash": _stable_hash(record),
            "semantic_hash": _stable_hash({
                "name": record.get("name") or record.get("canonical_name"),
                "record_type": record.get("record_type"),
            }),
            "disposition": disposition,
            "disposition_reason": f"carried from exercise audit disposition={record.get('disposition')}; migration remains blocked",
            "audit_confidence": "existing_audit_reconciled",
            "scientific_review_status": "pending_qualified_review",
            "human_review_required": True,
            "generator_eligible": False,
            "migration_eligible": False,
            "record": record,
        })
    return rows, {
        "present": True,
        "path": path.relative_to(ROOT).as_posix(),
        "record_count": len(rows),
        "source_summary": data.get("summary") or {},
        "reconciliation": data.get("reconciliation") or {},
        "graph_audit": data.get("graph_audit") or {},
    }


def load_workbook_records(workbook_path: Path) -> Tuple[List[Dict[str, Any]], List[Dict[str, Any]], Dict[str, Any]]:
    if not workbook_path.exists():
        return [], [], {"present": False, "path": str(workbook_path), "blocking_issue": "saved workbook not found"}
    try:
        import openpyxl
    except ImportError:
        return [], [], {"present": True, "path": str(workbook_path), "blocking_issue": "openpyxl unavailable"}

    wb = openpyxl.load_workbook(workbook_path, read_only=False, data_only=False)
    sheet_names = wb.sheetnames
    expected_final = {
        "Taxonomy", "Exercise Tags", "Exercise Relations", "Athletic Drills",
        "Training Protocols", "Team Templates", "Team Stations", "Media Registry",
        "Revision Log", "Validation Summary",
    }
    missing_final_sheets = sorted(expected_final.difference(sheet_names))
    rows: List[Dict[str, Any]] = []
    evidence: List[Dict[str, Any]] = []

    if "Curated Exercises" in wb.sheetnames:
        ws = wb["Curated Exercises"]
        headers = [str(cell.value or "").strip() for cell in ws[4]]
        for row_index in range(5, ws.max_row + 1):
            values = [ws.cell(row_index, column).value for column in range(1, ws.max_column + 1)]
            record = {headers[i]: values[i] for i in range(min(len(headers), len(values))) if headers[i]}
            if not any(value not in (None, "") for value in values):
                continue
            source_id = str(record.get("Review ID") or f"workbook_row_{row_index}")
            text = _record_text(record)
            approval_claim = any(term in text for term in ("approved", "generator eligible", "final coach/physio"))
            rows.append({
                "audit_record_id": f"ari_{_stable_hash([str(workbook_path), row_index, source_id])[:24]}",
                "source_id": source_id,
                "audit_only_id": None,
                "id_origin": "source",
                "source_path": str(workbook_path),
                "source_locator": f"'Curated Exercises'!{row_index}:{row_index}",
                "source_kind": "saved_workbook_snapshot",
                "record_type": "workbook_exercise_candidate",
                "sport": None,
                "in_launch_scope": True,
                "domain": str(record.get("Section") or "exercise"),
                "content_scope": "physical_preparation",
                "knowledge_layer": "layer_1_method_library",
                "runtime_collection_candidate": False,
                "evidence_linked": "Sources:" in str(record.get("Reviewer Notes") or ""),
                "evidence_reference_count": 0,
                "has_stable_source_id": True,
                "has_bounded_prescription": False,
                "claim_support_class": "evidence_linked_implication_candidate" if "Sources:" in str(record.get("Reviewer Notes") or "") else "unsupported_or_unlinked_assertion",
                "claim_support_classification_status": "saved_snapshot_pending_manual_verification",
                "payload_hash": _stable_hash(record),
                "semantic_hash": _stable_hash({"name": record.get("Exercise Name")}),
                "disposition": "quarantine" if approval_claim else "rewrite",
                "disposition_reason": (
                    "saved snapshot contains obsolete approval/generator claims and is missing final normalized sheets"
                    if approval_claim else "saved workbook snapshot requires reconciliation with the newer connected workbook"
                ),
                "audit_confidence": "snapshot_mismatch_detected",
                "scientific_review_status": "approval_claim_invalidated_pending_qualified_review",
                "human_review_required": True,
                "generator_eligible": False,
                "migration_eligible": False,
                "record": record,
            })

    if "Evidence Sources" in wb.sheetnames:
        ws = wb["Evidence Sources"]
        headers = [str(cell.value or "").strip() for cell in ws[1]]
        for row_index in range(2, ws.max_row + 1):
            values = [ws.cell(row_index, column).value for column in range(1, ws.max_column + 1)]
            if not any(value not in (None, "") for value in values):
                continue
            record = {headers[i]: values[i] for i in range(min(len(headers), len(values))) if headers[i]}
            evidence.append({
                "source_id": record.get("Source ID"),
                "title": record.get("Source"),
                "year": record.get("Year"),
                "source_type": record.get("Evidence Type"),
                "notes": record.get("Why It Matters"),
                "url": record.get("URL"),
                "source_path": str(workbook_path),
                "source_locator": f"'Evidence Sources'!{row_index}:{row_index}",
            })

    return rows, evidence, {
        "present": True,
        "path": str(workbook_path),
        "sheet_names": sheet_names,
        "sheet_count": len(sheet_names),
        "curated_row_count": len(rows),
        "evidence_source_count": len(evidence),
        "missing_expected_final_sheets": missing_final_sheets,
        "snapshot_status": "outdated_incomplete_snapshot" if missing_final_sheets else "final_structure_present",
        "blocking_issue": (
            "Save the latest connected workbook and rerun the audit before exercise migration decisions."
            if missing_final_sheets else None
        ),
    }


def _apply_duplicate_dispositions(rows: List[Dict[str, Any]]) -> Dict[str, Any]:
    id_groups: Dict[str, List[Dict[str, Any]]] = defaultdict(list)
    semantic_groups: Dict[Tuple[Any, ...], List[Dict[str, Any]]] = defaultdict(list)
    for row in rows:
        if row.get("source_id"):
            id_groups[str(row["source_id"])].append(row)
        semantic_groups[(row.get("sport"), row.get("record_type"), row.get("semantic_hash"))].append(row)

    id_collisions: List[Dict[str, Any]] = []
    exact_duplicate_groups: List[Dict[str, Any]] = []
    semantic_duplicate_groups: List[Dict[str, Any]] = []

    for source_id, group in sorted(id_groups.items()):
        if len(group) < 2:
            continue
        payload_hashes = {row["payload_hash"] for row in group}
        item = {
            "source_id": source_id,
            "record_count": len(group),
            "payload_variant_count": len(payload_hashes),
            "locations": [f"{row['source_path']}::{row['source_locator']}" for row in group],
        }
        if len(payload_hashes) == 1:
            exact_duplicate_groups.append(item)
        else:
            id_collisions.append(item)
            for row in group:
                row["disposition"] = "quarantine"
                row["disposition_reason"] = "duplicate source ID has conflicting payload variants; canonical resolution required"

    for key, group in semantic_groups.items():
        if len(group) < 2 or not key[2]:
            continue
        source_ids = {row.get("source_id") for row in group}
        if len(source_ids) < 2:
            continue
        semantic_duplicate_groups.append({
            "sport": key[0],
            "record_type": key[1],
            "semantic_hash": key[2],
            "source_ids": sorted(str(value) for value in source_ids if value),
            "locations": [f"{row['source_path']}::{row['source_locator']}" for row in group],
            "recommendation": "manual_merge_review",
        })
        preferred = sorted(
            group,
            key=lambda row: (
                bool(row.get("runtime_collection_candidate")),
                row.get("source_kind") == "compiled_collection",
                bool(row.get("has_stable_source_id")),
                bool(row.get("evidence_linked")),
                "draft" not in str(row.get("source_path", "")).lower(),
            ),
            reverse=True,
        )[0]
        for row in group:
            if row is preferred or row["disposition"] == "quarantine":
                continue
            row["disposition"] = "merge"
            row["disposition_reason"] = f"semantic duplicate of preferred candidate {preferred.get('source_id') or preferred.get('audit_only_id')}"

    return {
        "id_collision_count": len(id_collisions),
        "id_collisions": id_collisions,
        "exact_duplicate_group_count": len(exact_duplicate_groups),
        "exact_duplicate_groups": exact_duplicate_groups,
        "semantic_duplicate_group_count": len(semantic_duplicate_groups),
        "semantic_duplicate_groups": semantic_duplicate_groups,
    }


def _evidence_tier(source_type: Any, url: Any) -> int:
    text = f"{source_type or ''} {url or ''}".lower().replace("-", "_")
    for tier, hints in SOURCE_TIER_HINTS:
        if any(hint in text for hint in hints):
            return tier
    if "pubmed.ncbi.nlm.nih.gov" in text or "pmc.ncbi.nlm.nih.gov" in text:
        return 4
    if url:
        return 5
    return 0


def build_evidence_report(rows: Sequence[Dict[str, Any]], workbook_evidence: Sequence[Dict[str, Any]]) -> Dict[str, Any]:
    registry: Dict[str, Dict[str, Any]] = {}

    def add(ref: Mapping[str, Any], usage: Mapping[str, Any]) -> None:
        title = str(ref.get("title") or ref.get("Source") or ref.get("name") or "").strip()
        url = str(ref.get("url") or ref.get("URL") or "").strip()
        source_type = str(ref.get("source_type") or ref.get("Evidence Type") or "unspecified").strip()
        key = url or f"title:{_norm(title)}"
        if not key or key == "title:":
            key = f"unresolved:{_stable_hash([ref, usage])[:16]}"
        entry = registry.setdefault(key, {
            "source_key": key,
            "title": title or None,
            "url": url or None,
            "source_type": source_type,
            "evidence_tier": _evidence_tier(source_type, url),
            "usage_count": 0,
            "sports": set(),
            "record_types": set(),
            "locations": [],
            "online_support_verified": False,
            "verification_status": "provenance_only_not_semantically_verified",
        })
        entry["usage_count"] += 1
        if usage.get("sport"):
            entry["sports"].add(usage["sport"])
        if usage.get("record_type"):
            entry["record_types"].add(usage["record_type"])
        location = usage.get("location")
        if location and len(entry["locations"]) < 50:
            entry["locations"].append(location)

    for row in rows:
        for ref in _source_refs(row.get("record") or {}):
            add(ref, {
                "sport": row.get("sport"),
                "record_type": row.get("record_type"),
                "location": f"{row['source_path']}::{row['source_locator']}",
            })
    for ref in workbook_evidence:
        add(ref, {
            "sport": None,
            "record_type": "workbook_evidence_source",
            "location": f"{ref.get('source_path')}::{ref.get('source_locator')}",
        })

    sources: List[Dict[str, Any]] = []
    for entry in registry.values():
        item = dict(entry)
        item["sports"] = sorted(item["sports"])
        item["record_types"] = sorted(item["record_types"])
        sources.append(item)
    sources.sort(key=lambda item: (item["evidence_tier"] or 99, item.get("title") or item["source_key"]))
    tiers = Counter(item["evidence_tier"] for item in sources)
    claim_classes = Counter(row.get("claim_support_class") for row in rows)
    return {
        "audit_version": AUDIT_VERSION,
        "scope_note": "Source linkage and hierarchy audit only. Claim-to-source semantic support remains pending manual scientific review.",
        "source_count": len(sources),
        "sources_missing_url": sum(not item.get("url") for item in sources),
        "sources_by_tier": {str(key): value for key, value in sorted(tiers.items())},
        "record_claim_support_classes": dict(sorted(claim_classes.items())),
        "sources": sources,
    }


def build_coverage(rows: Sequence[Dict[str, Any]]) -> Dict[str, Any]:
    matrix: Dict[str, Any] = {}
    for sport in sorted(LAUNCH_SPORTS | DEFERRED_SPORTS):
        sport_rows = [
            row for row in rows
            if row.get("sport") == sport and row.get("content_scope") == "physical_preparation"
        ]
        dimension_rows: Dict[str, Any] = {}
        for dimension, terms in DIMENSIONS.items():
            matches = []
            for row in sport_rows:
                text = _substantive_text(row.get("record") or {})
                if any(term in text for term in terms):
                    matches.append(row)
            evidence_count = sum(bool(row.get("evidence_linked")) for row in matches)
            quantified_count = sum(bool(re.search(r"\b\d+(?:\.\d+)?\s*(?:%|s|sec|seconds|min|minutes|m|km|kg|week|weeks|reps|sets)\b", _substantive_text(row.get("record") or {}))) for row in matches)
            if not matches:
                status = "gap"
            elif len(matches) < 3 or not evidence_count:
                status = "weak"
            elif not quantified_count and dimension in {
                "force_requirements", "movement_speed", "effort_duration",
                "recovery_between_efforts", "repeated_effort_demands",
            }:
                status = "partial_unquantified"
            else:
                status = "candidate_coverage"
            dimension_rows[dimension] = {
                "status": status,
                "matching_record_count": len(matches),
                "evidence_linked_count": evidence_count,
                "quantified_record_count": quantified_count,
                "example_source_ids": [row.get("source_id") or row.get("audit_only_id") for row in matches[:5]],
                "human_scientific_review_required": True,
            }

        scopes = Counter(row.get("content_scope") for row in rows if row.get("sport") == sport)
        record_types = Counter(row.get("record_type") for row in sport_rows)
        statuses = Counter(value["status"] for value in dimension_rows.values())
        matrix[sport] = {
            "physical_preparation_record_count": len(sport_rows),
            "all_record_count": sum(scopes.values()),
            "content_scope_counts": dict(sorted(scopes.items())),
            "record_type_counts": dict(sorted(record_types.items())),
            "dimension_status_counts": dict(sorted(statuses.items())),
            "dimensions": dimension_rows,
            "release_status": "hidden_no_research_package" if sport == "hyrox" else "candidate_data_audit_pending",
        }
    return {
        "audit_version": AUDIT_VERSION,
        "method": "Keyword and structure coverage is a gap detector, not scientific approval.",
        "sports": matrix,
    }


def _write_json(path: Path, payload: Any) -> None:
    path.write_text(json.dumps(payload, indent=2, ensure_ascii=False, default=str) + "\n", encoding="utf-8")


def _write_jsonl(path: Path, rows: Iterable[Mapping[str, Any]], *, include_payload: bool) -> int:
    count = 0
    with path.open("w", encoding="utf-8") as handle:
        for row in rows:
            item = dict(row)
            if not include_payload:
                item.pop("record", None)
            handle.write(json.dumps(item, ensure_ascii=False, default=str, separators=(",", ":")) + "\n")
            count += 1
    return count


def _markdown_table(headers: Sequence[str], rows: Sequence[Sequence[Any]]) -> str:
    lines = ["| " + " | ".join(headers) + " |", "|" + "|".join("---" for _ in headers) + "|"]
    for row in rows:
        lines.append("| " + " | ".join(str(value).replace("|", "\\|").replace("\n", " ") for value in row) + " |")
    return "\n".join(lines)


def build_input_output_report() -> str:
    input_fields = [
        "sports", "sport_details", "primary_sport", "role_or_position", "event_or_discipline",
        "competition_level", "season_phase", "upcoming_events", "practice_schedule",
        "training_days_per_week", "preferred_training_days", "session_duration_min",
        "equipment", "facilities", "experience", "fitness_assessment", "pain_areas",
        "current_injuries", "stress_level", "sleep_avg_hours", "secondary_sports",
    ]
    collected_fields = {
        "sports", "sport_details", "competition_level", "season_phase",
        "training_days_per_week", "preferred_training_days", "session_duration_min",
        "equipment", "experience", "fitness_assessment", "pain_areas",
        "current_injuries", "stress_level",
    }
    modeled_fields = {
        "sports", "sport_details", "competition_level", "season_phase", "upcoming_events",
        "training_days_per_week", "preferred_training_days", "session_duration_min",
        "equipment", "facilities", "experience", "fitness_assessment", "pain_areas",
        "current_injuries", "stress_level",
    }
    consumers_by_field = {
        "sports": "macro_plan_service, knowledge_retrieval, ai_workout_service, server",
        "sport_details": "macro_plan_service, knowledge_retrieval, ai_workout_service",
        "role_or_position": "knowledge_retrieval via untyped sport_details only",
        "event_or_discipline": "no canonical consumer; occasional untyped sport_details inspection",
        "competition_level": "ai_workout_service",
        "season_phase": "macro_plan_service, knowledge_retrieval, ai_workout_service, server",
        "upcoming_events": "macro_plan_service, ai_workout_service",
        "training_days_per_week": "macro_plan_service, ai_workout_service, server",
        "preferred_training_days": "ai_workout_service, server",
        "session_duration_min": "macro_plan_service, ai_workout_service, server",
        "equipment": "knowledge_retrieval, ai_workout_service, server",
        "facilities": "knowledge_retrieval, ai_workout_service, server",
        "experience": "macro_plan_service, knowledge_retrieval, ai_workout_service, server",
        "fitness_assessment": "ai_workout_service, server",
        "pain_areas": "macro_plan_service, knowledge_retrieval, ai_workout_service, server",
        "current_injuries": "macro_plan_service, knowledge_retrieval, ai_workout_service, server",
        "stress_level": "macro_plan_service, ai_workout_service, server",
        "sleep_avg_hours": "macro_plan_service, ai_workout_service",
    }
    rows = []
    for field in input_fields:
        collected = field in collected_fields
        modeled = field in modeled_fields
        used_by = consumers_by_field.get(field, "—")
        if field == "sport_details":
            note = "Saved, but current onboarding populates only the sport name; role/event remains absent."
        elif not collected:
            note = "Required product input is missing from current onboarding."
        elif used_by == "—":
            note = "Collected/modelled but no meaningful planner consumption detected."
        else:
            note = "Present; semantic quality still requires implementation review."
        rows.append((field, "Yes" if collected else "No", "Yes" if modeled else "No", used_by, note))

    output_rows = [
        ("Program title/goal", "User", "Keep"),
        ("Athlete analysis", "User/internal", "Rewrite as structured deterministic needs summary"),
        ("Blocks and weeks", "User/internal", "Keep with template/rule versions"),
        ("Adaptation targets", "User", "Keep controlled quality codes plus labels"),
        ("Sport transfer", "User", "Keep short explanation; never use as selection authority"),
        ("Exercise IDs and prescriptions", "User/runtime", "Keep; IDs must resolve and prescriptions must be bounded"),
        ("Evidence/source objects", "Admin", "Do not send in ordinary workout responses"),
        ("Taxonomy and audit metadata", "Admin", "Internal only"),
        ("AI reasoning", "None", "Do not persist or expose"),
    ]
    return f"""# Input and Output Relevance Audit

## Athlete inputs

{_markdown_table(("Field", "Collected", "Backend model", "Detected consumers", "Finding"), rows)}

## Output classification

{_markdown_table(("Output", "Audience", "Decision"), output_rows)}

## Principal gaps

- There is no canonical primary sport or secondary-sport workload model.
- Position, running event, cycling discipline, swimming stroke/distance, competition dates, and external practice schedule are not collected with executable structure.
- `sport_details` is created as `{{sport}}` only, so existing role records cannot reliably affect planning.
- Current output validation emphasizes wording and shape; it does not prove macrocycle coherence, workload spacing, event suitability, or reproducibility.
"""


def build_data_flow_report() -> str:
    return """# Current Generator Data Flow and Audit Findings

```text
Onboarding store
  -> POST /onboarding/complete
  -> athlete_profiles.raw_profile + user profile
  -> FastAPI background generation
  -> macro_plan_service template scoring
  -> Mongo knowledge_retrieval queries
  -> compact_context_for_ai
  -> LLM prompt and JSON generation
  -> Pydantic parse
  -> validate_program_quality
  -> training_programs / program_blocks / workouts
  -> completion feedback / athlete_state
  -> another LLM-generated week
```

## Reproducibility finding

The current system is not reproducible from `athlete state + content versions + rule versions + template + deterministic decisions` because the LLM still chooses the program structure, exercise mix, dosage, and progression language. Content and rules lack a single versioned relational snapshot attached to every decision.

## Responsibilities to move out of the prompt

- Sport/event needs analysis.
- Macrocycle and phase selection.
- Weekly high/low scheduling.
- Competition-week placement.
- Exercise eligibility and substitution.
- Sets, reps, duration, intensity and recovery bounds.
- Progression, deload and missed-session handling.
- Pain, scope and external-load conflicts.

## Existing safeguards worth retaining

- Global generation kill switch.
- Exact exercise-ID grounding requirement.
- Schema parsing and retry ceiling.
- Per-user rate and quota guards.
- Atomic next-week generation claim.
- Persisted completion, RPE and pain feedback.

## Boundary

The generator must consume physical-preparation records only. Technical and tactical teaching records may remain a separate product library, but cannot enter workout selection or justify a physical prescription.
"""


def build_running_gap_report(rows: Sequence[Dict[str, Any]], coverage: Mapping[str, Any]) -> str:
    running_rows = [row for row in rows if row.get("sport") == "running_endurance"]
    event_terms = {
        "run_walk_fitness": ("run-walk", "run_walk", "fitness running"),
        "100m": ("100m", "100 m", "100-meter", "100 meter"),
        "200m": ("200m", "200 m", "200-meter", "200 meter"),
        "400m": ("400m", "400 m", "400-meter", "400 meter"),
        "5k": ("5k", "5 km"),
        "10k": ("10k", "10 km"),
        "half_marathon": ("half marathon", "21.1"),
        "marathon": ("marathon", "42.2"),
        "post_clearance_return": ("return to run", "return-to-run", "return_to_run"),
    }
    event_rows = []
    for event, terms in event_terms.items():
        matches = [row for row in running_rows if any(term in _substantive_text(row.get("record") or {}) for term in terms)]
        structured = sum(row.get("has_bounded_prescription", False) for row in matches)
        event_rows.append((event, len(matches), structured, "Gap" if not matches else "Candidate only — scientific review pending"))
    dimension_rows = coverage["sports"]["running_endurance"]["dimensions"]
    gap_dimensions = [name for name, item in dimension_rows.items() if item["status"] != "candidate_coverage"]
    return f"""# Running Reference Model Gap Report

## Current event coverage

{_markdown_table(("Event", "Candidate records", "Bounded records", "Status"), event_rows)}

## Structural findings

- The current running package is primarily endurance-oriented and contains 17 coarse workout archetypes expressed as text arrays.
- Sprint-event macrocycles for 100m, 200m and 400m are not represented as executable phase, weekly-load and prescription models.
- Road-event content does not provide one canonical event model for run-walk, 5K, 10K, half marathon and marathon across beginner/intermediate/advanced levels.
- Pace prescription lacks a deterministic hierarchy for recent race result, critical-speed/threshold estimate, RPE, talk test, heart rate and environmental adjustment.
- Strength support is described as a principle rather than integrated into event-specific weekly recipes.
- Return-to-running content must be limited to post-clearance re-entry and stop/modify rules; it cannot become rehabilitation treatment.
- Existing `trail_hill` content is outside Running v1 and should be quarantined until trail scope is approved.

## Demand dimensions requiring more or better evidence

{', '.join(gap_dimensions) if gap_dimensions else 'None detected by structural coverage; claim-level review is still mandatory.'}

## Required Running research packages

1. Sprint demand and planning: 100m, 200m, 400m.
2. Road/endurance demand and planning: run-walk, 5K, 10K, half marathon, marathon.
3. Event-specific strength, plyometric and tissue-capacity support.
4. Warm-up, sprint exposure, high/low scheduling, taper and competition-week rules.
5. Post-clearance return-to-running entry criteria and conservative stop rules.
6. Deterministic pace/intensity calculation and environmental adjustment.
7. Validated assessments and progression gates by event and training age.
"""


def build_architecture_report() -> str:
    return """# Cross-Sport Architecture Requirements

## Canonical knowledge layers

1. **Method library** — exercises, athletic drills, running/conditioning modalities, equipment, environments, instructions, relations, safety and media gates.
2. **Exercise effects** — sparse primary/secondary adaptation effects plus fatigue, impact, technical and supervision costs.
3. **Sport demands** — evidence-backed sport/event/role facts with value, unit, context, population, confidence and direct/inferred status.
4. **Prescriptions and constraints** — executable condition/action rules for frequency, dose, recovery, progression, deload, competition and conservative stopping.
5. **Program recipes** — program archetypes, phases, weekly structures, session blocks and selection slots.

## PostgreSQL-ready entities

- `knowledge_sources`, `evidence_claims`, `content_versions`, `content_reviews`.
- `sports`, `sport_events`, `sport_roles`, `physical_qualities`, `sport_demand_facts`, `sport_quality_priorities`.
- `training_methods`, `method_effects`, `method_constraints`, `method_relations`, `method_media`.
- `prescription_rules`, `progression_rules`, `competition_rules`, `external_load_conflicts`.
- `program_archetypes`, `program_phases`, `weekly_recipes`, `session_recipes`, `recipe_slots`.
- `athlete_sport_profiles`, `athlete_schedules`, `athlete_states`, `training_plans`, `plan_weeks`, `planned_sessions`, `prescribed_items`, `completion_records`, `adaptation_decisions`.

## Identifier and version rules

- One canonical UUID per production entity; preserve valid legacy UUIDs and create UUIDv7 for new entities.
- Human-readable codes and provider/source IDs are unique metadata, never competing primary IDs.
- Every generated plan stores the exact content, rule, template and algorithm versions used.
- All import and generation operations are idempotent.

## Planner boundary

```text
normalize athlete context
-> hard eligibility filters
-> sport/event needs vector
-> macrocycle and phase selection
-> weekly load placement
-> session recipe expansion
-> method selection
-> bounded prescription
-> invariant validation
-> versioned persistence
```

AI is not permitted to bypass these steps. A future AI layer may explain already-valid decisions or normalize free-text feedback into bounded structured signals.
"""


def build_migration_report(rows: Sequence[Dict[str, Any]], workbook_meta: Mapping[str, Any], duplicate_report: Mapping[str, Any]) -> str:
    dispositions = Counter(row["disposition"] for row in rows)
    layers = Counter(row["knowledge_layer"] for row in rows)
    table_rows = [(key, value) for key, value in sorted(dispositions.items())]
    layer_rows = [(key, value) for key, value in sorted(layers.items())]
    return f"""# Migration Eligibility Report

## Decision

**Migration is blocked.** No audited content record is generator-eligible or migration-eligible at this gate.

## Candidate dispositions

{_markdown_table(("Disposition", "Records"), table_rows)}

## Layer mapping

{_markdown_table(("Knowledge layer", "Records"), layer_rows)}

## Blocking conditions

- All automated dispositions require claim-level and domain review.
- The saved workbook status is `{workbook_meta.get('snapshot_status')}` and is missing {len(workbook_meta.get('missing_expected_final_sheets') or [])} expected final sheets.
- {duplicate_report.get('id_collision_count', 0)} duplicate-ID collision groups require canonical resolution.
- {duplicate_report.get('semantic_duplicate_group_count', 0)} semantic duplicate groups require manual merge decisions.
- Physical-preparation rules are mixed with technical/tactical and medical content in runtime candidate collections.
- Current prescriptions are not consistently structured or bounded.
- No local PostgreSQL schema, reviewed import manifest or cutover comparison exists yet.

## Gate to Task 2

Approve the audit dispositions and save the latest exercise workbook. Then define the canonical relational schema and import only records explicitly selected for rewrite/retention—not the current Mongo collections wholesale.
"""


def build_backlog() -> str:
    return """# Ordered Implementation Backlog

1. **Resolve Task-1 blockers** — save/reconcile the final workbook; review ID collisions, duplicate groups, scope leakage and evidence gaps.
2. **Canonical schemas** — define Pydantic/JSON Schema and PostgreSQL migrations for the five knowledge layers and runtime plans.
3. **Running research pack** — produce claim-level evidence records for 100m, 200m, 400m, run-walk, 5K, 10K, half marathon, marathon and post-clearance return.
4. **Running content normalization** — rewrite modalities, effects, prescription rules and recipes into original Runlete records with closed release gates.
5. **Onboarding v3** — collect primary/secondary sport, running event, experience/training age, target event/date, recent volume, practice schedule, equipment and conservative health context.
6. **Deterministic Running planner** — implement needs analysis, macrocycle, weekly placement, recipe expansion, prescription, validation and persistence without AI.
7. **Golden scenario validation** — cover every locked Running event, multi-sport conflicts, missed sessions, high RPE, pain flags and limited equipment.
8. **Bounded adaptation** — deterministic completion/readiness adjustments with audit history.
9. **Compatibility and local cutover** — expose versioned APIs and adapt existing program screens while Mongo remains authoritative.
10. **Sport expansion** — Football, Cricket, Basketball, Volleyball, Badminton, Tennis, Cycling, Swimming, Boxing, MMA, then Hyrox after its evidence gate.
11. **Optional AI layer** — explanations and feedback normalization only after deterministic plans pass all invariants.
"""


def build_summary(
    rows: Sequence[Dict[str, Any]],
    manifest: Mapping[str, Any],
    coverage: Mapping[str, Any],
    duplicate_report: Mapping[str, Any],
) -> str:
    dispositions = Counter(row["disposition"] for row in rows)
    scopes = Counter(row["content_scope"] for row in rows)
    sports = Counter(row["sport"] or "general_or_unassigned" for row in rows)
    summary_rows = [
        ("Inventory records", len(rows)),
        ("JSON files parsed", manifest["source_json"]["files_parsed"]),
        ("Saved workbook rows", manifest["workbook"].get("curated_row_count", 0)),
        ("Exercise-audit records", manifest["exercise_audit"].get("record_count", 0)),
        ("Generator eligible", 0),
        ("Migration eligible", 0),
        ("ID collision groups", duplicate_report.get("id_collision_count", 0)),
        ("Semantic duplicate groups", duplicate_report.get("semantic_duplicate_group_count", 0)),
    ]
    sport_rows = [(sport, count) for sport, count in sports.most_common()]
    coverage_rows = []
    for sport, item in coverage["sports"].items():
        statuses = item["dimension_status_counts"]
        coverage_rows.append((sport, item["physical_preparation_record_count"], statuses.get("candidate_coverage", 0), statuses.get("weak", 0), statuses.get("gap", 0), item["release_status"]))
    return f"""# Runlete Model Audit — Task 1 Summary

Generated: {manifest['generated_at']}

This audit is a migration and research decision aid. It grants no coaching, physiotherapy, medical, scientific or generator approval.

## Reconciliation

{_markdown_table(("Metric", "Count"), summary_rows)}

Every inventory record has exactly one disposition and every generator/migration gate remains closed.

## Dispositions

{_markdown_table(("Disposition", "Count"), sorted(dispositions.items()))}

## Content scope

{_markdown_table(("Scope", "Count"), sorted(scopes.items()))}

## Records by sport

{_markdown_table(("Sport", "Records"), sport_rows)}

## Structural coverage by sport

{_markdown_table(("Sport", "Physical records", "Candidate dimensions", "Weak dimensions", "Gap dimensions", "Release state"), coverage_rows)}

## Decision gate

Task 1 is complete when the generated outputs pass automated reconciliation and the product owner reviews the blockers. Task 2 must not begin by bulk-migrating the current Mongo collections. Running is the only sport authorized to proceed to a reference implementation after the audit gate.
"""


def run_audit(output_dir: Path = DEFAULT_OUTPUT, workbook_path: Path = DEFAULT_WORKBOOK) -> Dict[str, Any]:
    output_dir.mkdir(parents=True, exist_ok=True)
    raw_records, parse_errors, source_counts = discover_json_records()
    rows = [_inventory_row(raw) for raw in raw_records]
    exercise_rows, exercise_meta = load_exercise_audit_records()
    workbook_rows, workbook_evidence, workbook_meta = load_workbook_records(workbook_path)
    rows.extend(exercise_rows)
    rows.extend(workbook_rows)

    for row in rows:
        if row["disposition"] not in ALLOWED_DISPOSITIONS:
            raise ValueError(f"Invalid disposition {row['disposition']} at {row['source_path']}::{row['source_locator']}")
        row["generator_eligible"] = False
        row["migration_eligible"] = False

    duplicate_report = _apply_duplicate_dispositions(rows)
    unbounded = [
        {
            "source_id": row.get("source_id"),
            "sport": row.get("sport"),
            "record_type": row.get("record_type"),
            "location": f"{row['source_path']}::{row['source_locator']}",
        }
        for row in rows
        if row.get("knowledge_layer") in {"layer_4_prescriptions_constraints", "layer_5_program_recipes"}
        and row.get("content_scope") == "physical_preparation"
        and not row.get("has_bounded_prescription")
    ]
    scope_leakage = [
        {
            "source_id": row.get("source_id"),
            "sport": row.get("sport"),
            "record_type": row.get("record_type"),
            "content_scope": row.get("content_scope"),
            "location": f"{row['source_path']}::{row['source_locator']}",
        }
        for row in rows
        if row.get("runtime_collection_candidate")
        and row.get("content_scope") in {"technical_or_tactical", "medical_or_rehabilitation"}
    ]
    duplicate_report.update({
        "unbounded_prescription_candidate_count": len(unbounded),
        "unbounded_prescription_candidates": unbounded,
        "runtime_scope_leakage_count": len(scope_leakage),
        "runtime_scope_leakage": scope_leakage,
        "semantic_contradiction_status": "candidate_groups_only_manual_scientific_review_required",
    })

    coverage = build_coverage(rows)
    evidence = build_evidence_report(rows, workbook_evidence)
    generated_at = datetime.now(timezone.utc).isoformat()
    compact_exercise_meta = dict(exercise_meta)
    if compact_exercise_meta.get("reconciliation"):
        reconciliation = compact_exercise_meta["reconciliation"]
        compact_exercise_meta["reconciliation"] = {
            key: reconciliation.get(key)
            for key in (
                "candidate_rows", "collection_counts", "movement_candidates",
                "non_movement_candidates", "possible_equivalent_group_count",
                "unmatched_source_record_count",
            )
        }
    if compact_exercise_meta.get("graph_audit"):
        graph = compact_exercise_meta["graph_audit"]
        compact_exercise_meta["graph_audit"] = {
            key: graph.get(key)
            for key in (
                "edge_count", "progression_cycle_count", "self_edge_count",
                "unknown_reference_count",
            )
        }
    manifest: Dict[str, Any] = {
        "audit_version": AUDIT_VERSION,
        "generated_at": generated_at,
        "read_only_source_audit": True,
        "database_writes": False,
        "launch_sports": sorted(LAUNCH_SPORTS),
        "deferred_sports": sorted(DEFERRED_SPORTS),
        "source_json": {
            **source_counts,
            "parse_error_count": len(parse_errors),
            "parse_errors_accounted_as_artifacts": len(parse_errors),
            "unaccounted_parse_error_count": 0,
            "parse_errors": parse_errors,
        },
        "exercise_audit": compact_exercise_meta,
        "workbook": workbook_meta,
        "inventory_record_count": len(rows),
        "records_with_disposition": sum(row.get("disposition") in ALLOWED_DISPOSITIONS for row in rows),
        "generator_eligible_count": 0,
        "migration_eligible_count": 0,
        "required_outputs": list(REQUIRED_OUTPUTS),
    }

    _write_jsonl(output_dir / "source_inventory.jsonl", rows, include_payload=False)
    disposition_rows = [
        {
            key: row.get(key)
            for key in (
                "audit_record_id", "source_id", "audit_only_id", "source_path", "source_locator",
                "record_type", "sport", "content_scope", "knowledge_layer", "claim_support_class",
                "disposition", "disposition_reason", "scientific_review_status",
                "human_review_required", "generator_eligible", "migration_eligible",
            )
        }
        for row in rows
    ]
    _write_jsonl(output_dir / "record_dispositions.jsonl", disposition_rows, include_payload=False)
    _write_json(output_dir / "sport_coverage_matrix.json", coverage)
    _write_json(output_dir / "evidence_provenance.json", evidence)
    _write_json(output_dir / "duplicates_contradictions.json", duplicate_report)
    (output_dir / "current_generator_data_flow.md").write_text(build_data_flow_report(), encoding="utf-8")
    (output_dir / "input_output_relevance.md").write_text(build_input_output_report(), encoding="utf-8")
    (output_dir / "running_gap_report.md").write_text(build_running_gap_report(rows, coverage), encoding="utf-8")
    (output_dir / "cross_sport_architecture_requirements.md").write_text(build_architecture_report(), encoding="utf-8")
    (output_dir / "migration_eligibility.md").write_text(build_migration_report(rows, workbook_meta, duplicate_report), encoding="utf-8")
    (output_dir / "implementation_backlog.md").write_text(build_backlog(), encoding="utf-8")
    (output_dir / "audit_summary.md").write_text(build_summary(rows, manifest, coverage, duplicate_report), encoding="utf-8")
    _write_json(output_dir / "audit_manifest.json", manifest)
    return manifest


def validate_outputs(output_dir: Path) -> Dict[str, Any]:
    missing = [name for name in REQUIRED_OUTPUTS if not (output_dir / name).exists()]
    manifest = json.loads((output_dir / "audit_manifest.json").read_text(encoding="utf-8")) if not missing else {}
    inventory_count = 0
    disposition_count = 0
    invalid_dispositions: List[str] = []
    generator_open = 0
    migration_open = 0
    if (output_dir / "record_dispositions.jsonl").exists():
        with (output_dir / "record_dispositions.jsonl").open(encoding="utf-8") as handle:
            for line in handle:
                if not line.strip():
                    continue
                row = json.loads(line)
                inventory_count += 1
                if row.get("disposition") in ALLOWED_DISPOSITIONS:
                    disposition_count += 1
                else:
                    invalid_dispositions.append(str(row.get("disposition")))
                generator_open += row.get("generator_eligible") is not False
                migration_open += row.get("migration_eligible") is not False
    result = {
        "status": "pass",
        "missing_outputs": missing,
        "inventory_count": inventory_count,
        "manifest_inventory_count": manifest.get("inventory_record_count"),
        "records_with_valid_disposition": disposition_count,
        "invalid_dispositions": invalid_dispositions,
        "generator_gate_violations": generator_open,
        "migration_gate_violations": migration_open,
        "source_parse_errors": (manifest.get("source_json") or {}).get("parse_error_count"),
        "unaccounted_source_parse_errors": (manifest.get("source_json") or {}).get("unaccounted_parse_error_count"),
    }
    if (
        missing
        or invalid_dispositions
        or generator_open
        or migration_open
        or inventory_count != manifest.get("inventory_record_count")
        or disposition_count != inventory_count
        or (manifest.get("source_json") or {}).get("unaccounted_parse_error_count")
    ):
        result["status"] = "fail"
    return result


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Run the read-only Runlete model audit and generate Task-1 reports.")
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    parser.add_argument("--workbook", type=Path, default=DEFAULT_WORKBOOK)
    parser.add_argument("--validate-only", action="store_true")
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    if args.validate_only:
        result = validate_outputs(args.output)
    else:
        manifest = run_audit(args.output, args.workbook)
        result = validate_outputs(args.output)
        result["generated_at"] = manifest["generated_at"]
        result["output_dir"] = str(args.output)
    print(json.dumps(result, indent=2, ensure_ascii=False))
    if result["status"] != "pass":
        raise SystemExit(1)


if __name__ == "__main__":
    main()
