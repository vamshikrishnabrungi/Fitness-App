from __future__ import annotations

from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "Research materials" / "blueprints" / "Runlete_Workout_System_Rebuild_Reference.docx"

BLUE = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 77, 120)
INK = RGBColor(32, 38, 45)
MUTED = RGBColor(95, 105, 115)
LIGHT_GRAY = "F2F4F7"
LIGHT_BLUE = "E8EEF5"
LIGHT_GOLD = "FFF4CE"
LIGHT_RED = "FDEBEC"
WHITE = "FFFFFF"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_width(cell, width_dxa: int) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:type"), "dxa")
    tc_w.set(qn("w:w"), str(width_dxa))


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths: list[int]) -> None:
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:type"), "dxa")
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:type"), "dxa")
    tbl_ind.set(qn("w:w"), "120")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for index, cell in enumerate(row.cells):
            set_cell_width(cell, widths[index])
            set_cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def set_repeat_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    tr_pr.append(repeat)


def keep_with_next(paragraph) -> None:
    paragraph.paragraph_format.keep_with_next = True


def add_page_number(paragraph) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, separate, text, end])


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.font.color.rgb = INK
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    heading_tokens = {
        "Heading 1": (16, BLUE, 16, 8),
        "Heading 2": (13, BLUE, 12, 6),
        "Heading 3": (12, DARK_BLUE, 8, 4),
    }
    for name, (size, color, before, after) in heading_tokens.items():
        style = styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = color
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for name in ("List Bullet", "List Number"):
        style = styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(11)
        style.paragraph_format.left_indent = Inches(0.5)
        style.paragraph_format.first_line_indent = Inches(-0.25)
        style.paragraph_format.space_after = Pt(8)
        style.paragraph_format.line_spacing = 1.167

    header = section.header
    hp = header.paragraphs[0]
    hp.text = "RUNLETE  /  WORKOUT SYSTEM REBUILD"
    hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    hp.paragraph_format.space_after = Pt(2)
    hr = hp.runs[0]
    hr.font.name = "Calibri"
    hr._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    hr._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    hr.font.size = Pt(8.5)
    hr.font.bold = True
    hr.font.color.rgb = MUTED

    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    fr = fp.add_run("Internal rebuild reference  |  Page ")
    fr.font.name = "Calibri"
    fr._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    fr._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    fr.font.size = Pt(8.5)
    fr.font.color.rgb = MUTED
    add_page_number(fp)


def add_title_block(doc: Document) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run("WORKOUT SYSTEM REBUILD REFERENCE")
    r.font.name = "Calibri"
    r._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    r._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    r.font.size = Pt(23)
    r.font.bold = True
    r.font.color.rgb = INK

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(14)
    r = p.add_run("What to preserve from the Mongo prototype—and what to rebuild cleanly in PostgreSQL")
    r.font.size = Pt(13)
    r.font.color.rgb = MUTED

    rows = [
        ("Product", "Runlete physical-preparation planner"),
        ("Decision", "Clean PostgreSQL implementation; no wholesale Mongo content migration"),
        ("Reference sport", "Running v1 first"),
        ("Prepared", date.today().isoformat()),
        ("Status", "Filtered legacy-content purge completed and verified"),
    ]
    for label, value in rows:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        lr = p.add_run(f"{label}: ")
        lr.bold = True
        lr.font.color.rgb = INK
        vr = p.add_run(value)
        vr.font.color.rgb = INK


def add_callout(doc: Document, title: str, text: str, fill: str = LIGHT_BLUE) -> None:
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [9360])
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(title)
    r.bold = True
    r.font.color.rgb = DARK_BLUE
    p = cell.add_paragraph(text)
    p.paragraph_format.space_after = Pt(0)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


def add_numbered(doc: Document, items: list[str]) -> None:
    for number, item in enumerate(items, 1):
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.left_indent = Inches(0.375)
        paragraph.paragraph_format.first_line_indent = Inches(-0.188)
        paragraph.paragraph_format.space_after = Pt(4)
        paragraph.paragraph_format.line_spacing = 1.25
        paragraph.add_run(f"{number}.  {item}")


def add_table(doc: Document, headers: list[str], rows: list[list[str]], widths: list[int]):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    set_table_geometry(table, widths)
    header = table.rows[0]
    set_repeat_header(header)
    for i, text in enumerate(headers):
        cell = header.cells[i]
        set_cell_shading(cell, LIGHT_GRAY)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        r = p.add_run(text)
        r.bold = True
        r.font.color.rgb = INK
    for row_values in rows:
        cells = table.add_row().cells
        for i, text in enumerate(row_values):
            p = cells[i].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            r = p.add_run(str(text))
            r.font.size = Pt(9.5)
    set_table_geometry(table, widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def build() -> Path:
    doc = Document()
    configure_document(doc)
    add_title_block(doc)

    add_callout(
        doc,
        "Decision in one sentence",
        "The shared Mongo database was preserved. Only the 31 legacy workout-planning seed documents were removed; athlete-owned and running data remain intact, and the new knowledge and planner domains will be built directly in PostgreSQL.",
        LIGHT_GOLD,
    )

    doc.add_heading("1. What the live database actually contains", level=1)
    doc.add_paragraph(
        "The configured backend connects to one shared remote Mongo database named runlete. It is not divided into a workout database and a separate app database. At inspection time it contained 71 collections but only 35 documents in total."
    )
    add_table(
        doc,
        ["Data group", "Documents", "Decision"],
        [
            ["Legacy workout-planning seed content", "31 removed", "Verified at zero after filtered purge"],
            ["User account", "1", "Preserve"],
            ["Athlete profile", "1", "Preserve"],
            ["Recorded run", "1", "Preserve"],
            ["Club activity event", "1", "Preserve"],
            ["Saved workouts/programs/results", "0", "Nothing to migrate or delete"],
        ],
        [3900, 1200, 4260],
    )
    add_callout(
        doc,
        "Why the whole database must not be dropped",
        "Dropping runlete would erase the user, athlete profile, recorded run, club event, collection indexes and non-workout application structures. A content-scoped deletion achieves the clean rebuild without collateral data loss.",
        LIGHT_RED,
    )

    doc.add_heading("2. Exact legacy content removed", level=1)
    add_table(
        doc,
        ["Mongo collection", "Documents", "What it contains", "PostgreSQL replacement"],
        [
            ["macro_plan_templates", "6", "Coarse phase sequences and generic applicability rules", "Versioned program archetypes and phases"],
            ["planning_rules", "12", "Progression, readiness, pain and exercise-unlock candidates", "Executable prescription and constraint rules"],
            ["competition_week_rules", "7", "Generic match/race/bout-week structures", "Versioned competition and schedule rules"],
            ["sport_profiles", "6", "Broad Cricket, Volleyball, Football, Running, Basketball and combat summaries", "Evidence-backed sport/event/role demand facts"],
        ],
        [2050, 900, 3050, 3360],
    )
    doc.add_paragraph(
        "These 31 prototype seed documents were removed using a seed-source filter; no collection or database was dropped. Their concepts are summarized later in this brief, but their IDs, prose, prescriptions and evidence labels will not be copied automatically into the production model. The repository seed file remains a recoverable, non-authoritative reference."
    )

    doc.add_heading("3. Legacy workout collections that are currently empty", level=1)
    doc.add_paragraph(
        "The following Mongo collections contain zero documents. They need no data migration. They should be retired from schema creation and runtime dependencies only when their PostgreSQL replacements and API adapters are ready."
    )
    add_table(
        doc,
        ["Domain", "Empty collections"],
        [
            ["Generated plans and history", "macro_plans; athlete_states; training_programs; program_blocks; workouts; workout_sessions; exercise_results; user_exercise_history; user_level_assessments; user_benchmarks; ai_generation_log"],
            ["Method library", "exercise_library; primary_exercise_library; exercise_variation_library; exercise_progression_graph; mobility_drills; movement_patterns; physical_qualities; equipment_library"],
            ["Rules and templates", "training_protocols; workout_templates; progression_rules; readiness_rules; injury_modifications; benchmark_tests; running_workouts; running_plan_rules; programming_rules; recovery_rules"],
            ["Sport knowledge", "sport_roles; sport_training_rules; sport_teaching_progressions; sport_skill_assessments; sport_level_transition_rules; sport_library_progress"],
            ["Source-derived content", "knowledge_sources; source_registry; source_sections; knowledge_extraction_runs; training_principles; technical_models; technical_errors; coaching_progressions; glossary_terms"],
        ],
        [2200, 7160],
    )

    doc.add_heading("4. Data that must be preserved", level=1)
    add_bullets(
        doc,
        [
            "Identity and profile data: users and athlete_profiles.",
            "Running and competition data: terra_runs, activities, activity processing, routes, segments, territory, clubs, memberships, challenges, races and leaderboard data.",
            "Athlete health and self-report data: injuries, injury_logs, health_metrics, quick_logs, moods and daily snapshots—even where currently empty—because these are user-owned runtime domains, not knowledge content.",
            "Nutrition and food-analysis results: meals, nutrition_targets, nutrition_guidelines and related records; they are outside the workout-knowledge cleanup.",
            "Privacy, safety, account lifecycle, imports, integrations and notifications.",
            "Repository research and audit files as non-authoritative reference material. They should remain outside production retrieval until rewritten and approved under the new governance model.",
        ],
    )

    doc.add_heading("5. Product decisions worth preserving", level=1)
    add_bullets(
        doc,
        [
            "Runlete is a constrained-AI physical-preparation planner supported by reviewed knowledge: AI selects and composes only from approved template slots, candidate methods and dose bounds; deterministic validation remains authoritative.",
            "Running is the first complete reference implementation. Launch scope: run-walk/fitness, 5K, 10K, half marathon, marathon, 100m, 200m and 400m. Trail, ultra and 800m–1500m remain deferred.",
            "The broader launch set is Badminton, Basketball, Boxing, Cricket, Cycling, Football, MMA, Running, Swimming, Tennis and Volleyball. Hyrox remains hidden until researched and audited.",
            "An athlete has one primary sport. Secondary sports modify total workload, conflicts and supporting priorities rather than creating independent overlapping programs.",
            "Physical preparation is in scope; sport technique, tactics and ordinary rehabilitation treatment are not.",
            "Pain or injury inputs produce conservative modification, stopping guidance or referral. They do not cause the system to prescribe rehabilitation.",
            "Every accepted plan must be reproducible from its persisted input snapshot, content/rule/recipe versions, candidate packet, model output and validator version; idempotent retries reuse the accepted result.",
        ],
    )

    doc.add_heading("6. The production five-layer knowledge model", level=1)
    add_table(
        doc,
        ["Layer", "Purpose", "Examples"],
        [
            ["1. Method library", "Defines what can be prescribed once, independent of sport", "Exercises, drills, running modalities, equipment, instructions, safety, media, substitutions"],
            ["2. Method effects", "Connects a method to intended adaptations and costs", "Primary effect; up to three secondary effects; fatigue, impact, technical and supervision cost"],
            ["3. Sport demands", "Represents evidence-backed event and role demands—not exercise lists", "Force, direction, contraction, duration, recovery, tissue load, schedule, surface and environment"],
            ["4. Prescription rules", "Stores executable constraints and dose bounds", "Frequency, intensity, volume, recovery, spacing, progression, deload, pain and external-load conflicts"],
            ["5. Program recipes", "Defines reviewed structures into which eligible methods are selected", "Program archetypes, phases, weeks, sessions, blocks, slots and substitution policies"],
        ],
        [1900, 3360, 4100],
    )
    doc.add_paragraph(
        "A method exists once. Sports do not receive duplicated exercise libraries. A sport/event demand model defines required qualities; recipes and rules translate those qualities into sessions; the method library supplies eligible options."
    )

    doc.add_heading("7. Sport-demand information the new system must understand", level=1)
    add_bullets(
        doc,
        [
            "Force magnitude, impulse, rate of force development and power requirements.",
            "Dominant force directions: vertical, horizontal, lateral and rotational.",
            "Contraction behavior: concentric, eccentric, isometric and stretch-shortening actions.",
            "Movement velocity, contact time, effort duration and recovery distribution.",
            "Joint positions and meaningful ranges of motion under relevant loads.",
            "Repeated-effort demands, fatigue behavior and energy-system contributions.",
            "Competition format, weekly schedule, congestion, tapering and seasonal phase.",
            "Commonly stressed tissues expressed as load considerations—not injury-prevention promises.",
            "Event, position, level, age, sex, environment, equipment, surface and climate context.",
            "Monitoring measures and assessments that can actually change a programming decision.",
        ],
    )

    doc.add_heading("8. Minimum athlete input contract", level=1)
    add_table(
        doc,
        ["Input group", "Required structured information"],
        [
            ["Sport identity", "Primary sport; event/discipline; role/position where applicable; competition level; secondary sports"],
            ["Goal and calendar", "Goal type; priority event and date; season phase; practice, match and travel schedule"],
            ["Training availability", "Training days; preferred days; session duration; facilities; equipment; environment"],
            ["Training state", "Training age; recent training volume; recent event/performance result; assessments; completion history"],
            ["Recovery and constraints", "Sleep; stress; soreness; pain areas; current professional restrictions; recent layoff"],
            ["Sport-specific load", "Running volume and pace anchors; jump count; bowling/throwing load; sparring; pool/bike sessions, as relevant"],
        ],
        [2550, 6810],
    )
    doc.add_paragraph(
        "The current onboarding does not reliably collect primary sport, secondary-sport workload, role, event/discipline, practice schedule, upcoming competition or average sleep. The new planner cannot be considered complete until those inputs are typed, persisted and consumed."
    )

    doc.add_heading("9. Constrained-AI planner and deterministic validation boundary", level=1)
    add_numbered(
        doc,
        [
            "Normalize and validate athlete context.",
            "Apply hard eligibility and safety filters.",
            "Select the primary sport, event and goal demand model.",
            "Build a needs vector and resolve secondary-sport conflicts.",
            "Select the macrocycle archetype and phase structure.",
            "Place weekly high- and low-load demands around practices and competition.",
            "Expand reviewed session recipes and block requirements.",
            "Project only eligible released method IDs into each reviewed recipe slot.",
            "Ask AI to select and compose from that exact candidate set and choose only bounded prescription values.",
            "Validate schedule, dose, equipment, level, pain and workload invariants.",
            "Persist the plan plus input hash, candidate/output snapshot, provider, model, prompt, response schema, content, rule, recipe and validator versions.",
        ],
    )
    add_callout(
        doc,
        "AI boundary",
        "AI is allowed to choose among the exact eligible method IDs and dose bounds supplied for reviewed template slots. It may not invent IDs or fields, access draft/unrelated content, bypass safety and workload rules, diagnose injury or overwrite deterministic facts. Invalid output receives one structured repair attempt and then fails safely or uses a reviewed fallback.",
    )

    doc.add_heading("10. Existing implementation patterns worth reusing", level=1)
    add_bullets(
        doc,
        [
            "Global generation kill switch; it is currently disabled in the configured environment and should remain disabled through the rebuild.",
            "Exact approved method-ID grounding for every prescribed item.",
            "Typed response parsing, bounded retries and explicit failure states.",
            "Per-user rate and quota protection.",
            "Atomic claims for next-week generation or adaptation jobs.",
            "Persisted completion, session RPE, pain and exercise history as adaptation inputs.",
            "Separation of administrative evidence/taxonomy from athlete-facing workout responses.",
        ],
    )

    doc.add_heading("11. Useful prototype concepts—but not approved production rules", level=1)
    doc.add_paragraph(
        "The 31 populated seed documents contain several sensible planning ideas. They should be retained as research questions and design hypotheses, then rewritten into bounded rules only after evidence and expert review."
    )
    add_table(
        doc,
        ["Candidate concept", "What must be established before implementation"],
        [
            ["Progress one loading variable at a time", "Define the eligible variables, event context, rate limits, exceptions and failure conditions"],
            ["Place competition and practice before gym loading", "Model exact schedule conflicts, role exposure, congestion and recovery windows"],
            ["Keep power and sprint work high-quality and low-fatigue", "Define intensity, volume, rest, technical-quality and stopping thresholds by level"],
            ["Reduce volume while maintaining selected intensity during taper", "Create event-, athlete- and duration-specific taper bounds rather than one universal percentage"],
            ["Differentiate starters/high-minute athletes from underloaded athletes", "Require reliable external-load input and define safe top-up rules"],
            ["Use conservative pain and 24-hour-response gates", "Define non-diagnostic questions, stop rules, referral conditions and prohibited claims"],
            ["Do not stack redundant high-load exposures", "Quantify sprint, jump, eccentric, throwing, bowling, sparring and endurance conflicts"],
        ],
        [4000, 5360],
    )

    doc.add_heading("12. Running-first research and build scope", level=1)
    add_table(
        doc,
        ["Running event", "Current candidate status", "Required action"],
        [
            ["Run-walk / fitness", "Candidate material exists", "Normalize progression, intensity hierarchy and entry/stop rules"],
            ["5K", "Candidate material exists", "Build demand facts, phases, weekly recipes and pace hierarchy"],
            ["10K", "Candidate material exists", "Build demand facts, phases, weekly recipes and pace hierarchy"],
            ["Half marathon", "Material gap", "Research and author complete model"],
            ["Marathon", "Candidate material exists", "Complete long-run, fueling-interface, taper and load rules"],
            ["100m", "Material gap", "Research acceleration, maximum velocity, speed endurance and strength support"],
            ["200m", "Material gap", "Research sprint and speed-endurance integration"],
            ["400m", "Material gap", "Research speed, special endurance and recovery structure"],
            ["Post-clearance return", "Partial candidate material", "Restrict to conservative re-entry after professional clearance"],
        ],
        [1900, 2500, 4960],
    )
    doc.add_paragraph(
        "The present running package is endurance-biased and contains 17 coarse, prose-based workout archetypes. It does not yet provide executable sprint macrocycles, one canonical event model across levels, a deterministic pace hierarchy, integrated strength support or reliable secondary-sport conflict handling."
    )

    doc.add_heading("13. PostgreSQL ownership model", level=1)
    add_table(
        doc,
        ["Domain", "Representative PostgreSQL entities"],
        [
            ["Evidence and governance", "knowledge_sources, evidence_claims, content_versions, content_reviews, audit_events"],
            ["Taxonomy and demands", "sports, sport_events, sport_roles, physical_qualities, sport_demand_facts, sport_quality_priorities"],
            ["Methods", "training_methods, method_effects, method_constraints, method_relations, method_media"],
            ["Rules", "prescription_rules, progression_rules, competition_rules, external_load_conflicts"],
            ["Recipes", "program_archetypes, program_phases, weekly_recipes, session_recipes, recipe_slots"],
            ["Athlete runtime", "athlete_sport_profiles, athlete_schedules, athlete_states, training_plans, plan_weeks, planned_sessions, prescribed_items, completion_records, adaptation_decisions"],
        ],
        [2500, 6860],
    )
    add_bullets(
        doc,
        [
            "Use one canonical UUID per production entity and one unique stable code for human-readable references.",
            "New entities use UUIDv7; Mongo _id values never become public production identifiers.",
            "Foreign keys—not duplicated string arrays—connect demands, effects, rules, recipes and prescribed items.",
            "Every content record has draft, reviewed, approved, retired and superseded lifecycle states with effective dates.",
            "No record is generator-eligible merely because it was imported or inserted.",
            "Every plan is immutable by version; adaptations create explicit decisions rather than silently rewriting history.",
        ],
    )

    doc.add_heading("14. Clean rebuild sequence", level=1)
    add_numbered(
        doc,
        [
            "Completed: removed only the 31 populated legacy seed documents and preserved the four non-workout documents and shared database.",
            "Stop legacy workout collections from being automatically recreated or reseeded by db_setup and seed_all once the PostgreSQL replacement branch is ready.",
            "Create the PostgreSQL schemas and migrations for governance, taxonomy, methods, effects, demands, rules, recipes and runtime plans.",
            "Author the Running evidence package at claim level. Do not import the old prose as production truth.",
            "Create approved Running methods, effects, demand facts, rules and recipes using original Runlete wording.",
            "Implement the constrained-AI Running planner, strict selection schema, one repair attempt, immutable generation audit and invariant validator.",
            "Update onboarding and API contracts, then adapt the mobile workout UI to the new outputs.",
            "Validate golden scenarios before enabling generation.",
            "Expand sport by sport only after Running passes its release gate.",
        ],
    )

    doc.add_heading("15. Release gates", level=1)
    add_bullets(
        doc,
        [
            "No duplicate canonical identifiers or stable codes.",
            "No orphan evidence, method, rule, recipe or prescribed-item references.",
            "No technical/tactical sport instruction entering physical-preparation selection.",
            "No rehabilitation prescription generated from ordinary pain input.",
            "Every rule has executable conditions, bounded actions, scope and evidence provenance.",
            "Every plan fits athlete schedule, equipment, environment, level, event and competition calendar.",
            "Primary and secondary sports cannot create conflicting high-load sessions.",
            "Accepted plans store the full reproducibility lineage; duplicate requests are idempotent and do not trigger a second selection.",
            "No plan can be persisted or shown until all invariants pass.",
        ],
    )

    doc.add_heading("16. Deletion execution record", level=1)
    add_callout(
        doc,
        "Completed safe operation",
        "Deleted 31 documents from macro_plan_templates, planning_rules, competition_week_rules and sport_profiles where seed_source matched sftc_backend_planning_collections.json. The runlete database was not dropped. users, athlete_profiles, terra_runs and club_activity_events were preserved. Repository research and audit files were not deleted.",
        LIGHT_GOLD,
    )
    doc.add_paragraph(
        "A second, later code-removal step should retire Mongo workout schema creation, seeding, retrieval and generation paths after PostgreSQL endpoints exist. Performing that code removal now would break current workout-facing API routes before their replacement is available."
    )

    doc.add_page_break()
    doc.add_heading("Appendix A — Terms", level=1)
    terms = [
        ("Method", "A prescribable exercise, drill or conditioning modality."),
        ("Effect", "An intended adaptation plus the method's fatigue, impact and technical costs."),
        ("Demand fact", "An evidence-backed statement about a sport, event or role in a defined context."),
        ("Prescription rule", "Executable conditions and bounded actions that determine dose, recovery or eligibility."),
        ("Recipe", "A reviewed structural pattern for a program, week, session or block."),
        ("Generator eligible", "Explicitly approved for inclusion in a constrained AI candidate set; insertion into a database is not approval."),
        ("Post-clearance return", "Conservative re-entry after appropriate professional clearance, not treatment or rehabilitation."),
        ("External load", "Sport practice, competition or other training stress that the planner did not prescribe but must account for."),
    ]
    add_table(doc, ["Term", "Meaning"], [[a, b] for a, b in terms], [2500, 6860])

    doc.add_heading("Appendix B — Source basis for this brief", level=1)
    doc.add_paragraph(
        "This brief consolidates the read-only Task 1 audit, the live Mongo collection/count inspection, the current backend generator flow, the Running gap report and the agreed Runlete product constraints. It is a rebuild specification, not scientific approval of the prototype records."
    )
    add_bullets(
        doc,
        [
            "Research materials/audits/model/audit_summary.md",
            "Research materials/audits/model/current_generator_data_flow.md",
            "Research materials/audits/model/input_output_relevance.md",
            "Research materials/audits/model/running_gap_report.md",
            "Research materials/audits/model/cross_sport_architecture_requirements.md",
            "Research materials/audits/model/migration_eligibility.md",
            "Research materials/audits/model/implementation_backlog.md",
            "backend/db_setup.py and the configured read-only Mongo collection inspection",
        ],
    )

    for paragraph in doc.paragraphs:
        for run in paragraph.runs:
            if run.font.name is None:
                run.font.name = "Calibri"
                run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Calibri")
                run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Calibri")

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    print(build())
