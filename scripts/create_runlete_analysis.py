from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pathlib import Path

OUT = Path("/Users/vamshikrishna/Downloads/Fitness-App-main/docs/Runlete analysis.docx")
OUT.parent.mkdir(parents=True, exist_ok=True)
NAVY, TEAL, MUTED, LIGHT, PALE, BORDER = "102A43", "0B6E69", "52606D", "F4F7F6", "E8F5E9", "D9E2EC"

def shade(cell, fill):
    p = cell._tc.get_or_add_tcPr(); x = p.find(qn("w:shd"))
    if x is None: x = OxmlElement("w:shd"); p.append(x)
    x.set(qn("w:fill"), fill)

def margins(cell):
    p = cell._tc.get_or_add_tcPr(); x = p.first_child_found_in("w:tcMar")
    if x is None: x = OxmlElement("w:tcMar"); p.append(x)
    for side, val in (("top",100),("start",120),("bottom",100),("end",120)):
        n = x.find(qn("w:"+side))
        if n is None: n = OxmlElement("w:"+side); x.append(n)
        n.set(qn("w:w"), str(val)); n.set(qn("w:type"), "dxa")

def font(run, size=10.5, color="243B53", bold=False, italic=False, name="Aptos"):
    run.font.name=name; run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"),name); run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"),name)
    run.font.size=Pt(size); run.font.color.rgb=RGBColor.from_string(color); run.bold=bold; run.italic=italic

def ptext(doc, text="", size=10.5, color="243B53", bold=False, after=6, before=0, align=None):
    p=doc.add_paragraph(); p.paragraph_format.space_before=Pt(before); p.paragraph_format.space_after=Pt(after); p.paragraph_format.line_spacing=1.12
    if align is not None: p.alignment=align
    if text: font(p.add_run(text),size,color,bold)
    return p

def heading(doc,text,level=1):
    p=doc.add_paragraph(style="Heading "+str(level)); p.paragraph_format.keep_with_next=True
    font(p.add_run(text),{1:17,2:13,3:11.5}[level],{1:NAVY,2:TEAL,3:MUTED}[level],True)
    return p

def bullet(doc,text,numbered=False):
    p=doc.add_paragraph(style="List Number" if numbered else "List Bullet"); p.paragraph_format.space_after=Pt(3); p.paragraph_format.line_spacing=1.08
    font(p.add_run(text),10.2); return p

def callout(doc,label,text,fill=PALE,color=TEAL):
    t=doc.add_table(rows=1,cols=1); t.alignment=WD_TABLE_ALIGNMENT.LEFT; t.autofit=False
    c=t.cell(0,0); shade(c,fill); margins(c); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p=c.paragraphs[0]; p.paragraph_format.space_after=Pt(0); p.paragraph_format.line_spacing=1.1
    font(p.add_run(label.upper()+"  "),9,color,True); font(p.add_run(text),10.2)
    doc.add_paragraph().paragraph_format.space_after=Pt(1)

def table(doc,headers,rows):
    t=doc.add_table(rows=1,cols=len(headers)); t.alignment=WD_TABLE_ALIGNMENT.LEFT; t.autofit=False
    for i,h in enumerate(headers):
        c=t.rows[0].cells[i]; shade(c,NAVY); margins(c); font(c.paragraphs[0].add_run(h),9,"FFFFFF",True)
    for ri,row in enumerate(rows):
        cs=t.add_row().cells
        for i,val in enumerate(row):
            c=cs[i]; shade(c,"FFFFFF" if ri%2==0 else LIGHT); margins(c); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            q=c.paragraphs[0]; q.paragraph_format.space_after=Pt(0); q.paragraph_format.line_spacing=1.05; font(q.add_run(str(val)),9.1)
    return t

doc=Document(); sec=doc.sections[0]; sec.top_margin=Inches(.78); sec.bottom_margin=Inches(.75); sec.left_margin=Inches(.85); sec.right_margin=Inches(.85); sec.header_distance=Inches(.35); sec.footer_distance=Inches(.35)
normal=doc.styles["Normal"]; normal.font.name="Aptos"; normal._element.rPr.rFonts.set(qn("w:ascii"),"Aptos"); normal._element.rPr.rFonts.set(qn("w:hAnsi"),"Aptos"); normal.font.size=Pt(10.5); normal.font.color.rgb=RGBColor.from_string("243B53")
for n,s,c,b,a in [("Heading 1",17,NAVY,16,7),("Heading 2",13,TEAL,12,5),("Heading 3",11.5,MUTED,9,4)]:
    st=doc.styles[n]; st.font.name="Aptos Display"; st._element.rPr.rFonts.set(qn("w:ascii"),"Aptos Display"); st._element.rPr.rFonts.set(qn("w:hAnsi"),"Aptos Display"); st.font.size=Pt(s); st.font.bold=True; st.font.color.rgb=RGBColor.from_string(c); st.paragraph_format.space_before=Pt(b); st.paragraph_format.space_after=Pt(a); st.paragraph_format.keep_with_next=True
for n in ["List Bullet","List Number"]:
    doc.styles[n].font.name="Aptos"; doc.styles[n].font.size=Pt(10.2); doc.styles[n].paragraph_format.space_after=Pt(3)
hp=sec.header.paragraphs[0]; hp.alignment=WD_ALIGN_PARAGRAPH.RIGHT; font(hp.add_run("RUNLETE  /  PRODUCT & ENGINEERING REFERENCE"),8.5,MUTED,True)
fp=sec.footer.paragraphs[0]; fp.alignment=WD_ALIGN_PARAGRAPH.CENTER; font(fp.add_run("Runlete analysis  •  Internal reference  •  19 August 2026"),8,MUTED)
ptext(doc,"RUNLETE",11,TEAL,True,2,10); ptext(doc,"Runlete analysis",28,NAVY,True,4); ptext(doc,"A single guide to the product, running platform, competition system and deterministic training engine",13,MUTED,False,14)
callout(doc,"Purpose","This guide is the orientation layer for a new engineer, coach, operator or partner. It explains what Runlete is, how its parts connect, what is authoritative, and what remains planned versus implemented.","EAF4F4",TEAL)

heading(doc,"1. Runlete in one sentence")
ptext(doc,"Runlete is a running-first athletic platform that records and analyzes activities, verifies routes and road territory, powers competitive clubs, and generates sport-aware training programs from a reviewed exercise and template library.")
ptext(doc,"It combines the utility of a running tracker with adaptive physical preparation. It is not a social feed, messaging app, rehabilitation service, or general sport-technique coach.",after=8)
table(doc,["Product area","What Runlete does","What it deliberately does not do"],[
["Run tracking","Records GPS activities, calculates distance/time/pace and creates activity history.","Treat raw GPS drawings as verified roads or make frontend calculations authoritative."],
["Competition","Runs clubs, territory, challenges, races and leaderboards.","Posts, comments, kudos, follows, chat or direct messaging."],
["Training","Builds validated multi-week programs and selects workouts from approved methods.","Let an LLM invent exercises, IDs, unsafe doses or rehabilitation treatment."],
["Nutrition/health","Stores food estimates and readiness/pain context with user confirmation and conservative rules.","Diagnose illness or injury."],
["Admin Studio","Lets the operator manage approved methods, templates, programs and sport data.","Edit arbitrary backend code through the UI."]])

heading(doc,"2. The product mental model")
ptext(doc,"Every important user action follows the same pattern: capture an input, normalize it, apply reviewed knowledge and deterministic validation, persist the result, then show a concise athlete-facing output. AI is an assistant inside that pipeline; it is not the authority.")
for x in ["Athlete state: sport, event/role, goal, level, schedule, equipment, recent load, readiness and restrictions.","Knowledge: approved exercise methods, sport-demand priorities, four-week templates, dosage bounds, progressions, substitutions and safety rules.","Planner: uses athlete state and relevant knowledge to produce a four-week structure and selected sessions.","Validation: checks duration, hard/easy spacing, equipment, level, workload, pain rules and allowed IDs.","Delivery: mobile shows only the workout, instructions, dose, recovery and short rationale.","Feedback: completion, RPE, readiness and pain signals adjust future sessions within bounded rules."]: bullet(doc,x)

heading(doc,"3. Backend architecture")
ptext(doc,"The backend is a modular FastAPI application. PostgreSQL/PostGIS is the authoritative store for the current production domains. MongoDB and old local knowledge files are not runtime authorities in the rebuilt design.")
table(doc,["Layer","Technology / authority","Responsibility"],[
["API","FastAPI, Pydantic, SQLAlchemy 2","Versioned /api/v1 contracts, validation, authentication, policies and transactions."],
["Database","PostgreSQL 16 + PostGIS","Users, athletes, knowledge, plans, activities, clubs, territory, segments, races and leaderboards."],
["Object storage","Google Cloud Storage","Raw GPS streams, imports, food images, exports, exercise media and map graph packages."],
["Async events","PostgreSQL outbox + Pub/Sub + Cloud Run workers","Activity processing, map matching, territory recomputation, notifications, nutrition analysis and maintenance."],
["Cache/locks","Redis / Memorystore","Ephemeral cache, rate limits, locks and live standings; never the source of truth."],
["Maps","Mapbox rendering + self-hosted Valhalla matching","Map display and routing are separate from Runlete-owned verified edge and territory data."],
["AI","OpenAI gpt-4o currently configured","Food-image analysis and bounded workout selection/explanation; outputs are validated before persistence."],
["Operations","Cloud Run, Cloud SQL, Storage, Pub/Sub, Secret Manager","GCP deployment, backups, IAM, alerts, cost controls and releases."]])

heading(doc,"4. Database ownership and data flow")
ptext(doc,"Each domain owns its tables and service logic. New entities use application-generated UUIDv7 IDs. Stable codes and provider IDs are metadata with unique constraints, never alternate primary keys.")
callout(doc,"Authority rule","Raw evidence is immutable; cleaned streams, matched geometry and calculated facts are versioned derivatives. A newer algorithm may reprocess history, but it must never silently overwrite the original evidence.","FFF8E1","7A5A00")
table(doc,["Domain","Representative data","Key rule"],[
["Identity / athlete","Users, sessions, sport choices, schedule, equipment, readiness","Sensitive health fields are restricted and encrypted; APIs return only required fields."],
["Knowledge","Methods, effects, sport demands, rules, recipes, evidence, releases","Published releases are immutable; plans pin the exact release and versions used."],
["Training","Programs, phases, weeks, sessions, prescriptions, completion, adaptation","No invalid plan is persisted; future weeks can be regenerated against a newer release only explicitly."],
["Activity","Uploads, raw/clean streams, facts, routes, quality, segments, efforts","Processing is asynchronous and idempotent; frontend values are never authoritative."],
["Competition","Clubs, memberships, primary-club attribution, territory scores, challenges, races","An eligible activity contributes to only one primary club at activity time."],
["Operations","Outbox events, jobs, idempotency keys, audit logs, feature flags","Every retryable mutation has an idempotency key and stable error response."]])

heading(doc,"5. Running: recording an activity")
ptext(doc,"The recorder is a recoverable state machine rather than a single foreground request:")
callout(doc,"State machine","idle → recording → paused → finishing → uploaded → processing → provisional → complete / rejected","EAF4F4",TEAL)
ptext(doc,"On the device, GPS samples are persisted incrementally and queued offline. The client uploads resumable chunks with idempotency keys, so app backgrounding, termination, poor connectivity or retry does not create duplicate activities.")
heading(doc,"What is calculated server-side",2)
for x in ["Elapsed, moving and paused time; distance and corrected elevation.","Kilometre/mile splits, pace, grade-adjusted pace, heart-rate and cadence summaries when available.","Calories as a server-side estimate using approved inputs—not distance times a frontend constant.","GPS quality, duplicate timestamps, impossible jumps, vehicle-like speed and map-matching confidence.","Best efforts, personal records, route geometry and competition eligibility."]: bullet(doc,x)
heading(doc,"Activity quality and privacy",2)
ptext(doc,"Activities are public, club-visible or private. Hidden start/end zones are removed before maps, tiles, exports, timelines or moderator previews. Private activities never enter territory, races, challenges or leaderboards. Users aged 16–17 receive stricter defaults.")

heading(doc,"6. Maps, route detection and accurate analysis")
ptext(doc,"Runlete does not claim that a GPS line itself is a road. GPS is noisy: a runner can drift across a parallel road, pass under a flyover, lose signal in a tunnel or record a curved path that does not align with a mapped way. Verification therefore uses distinct layers.")
table(doc,["Layer","Purpose","Runlete behavior"],[
["Map display","Show streets, terrain and labels.","Mapbox mobile/vector rendering; it is a visual provider, not the owner of territory facts."],
["Map matching","Translate cleaned GPS samples to likely accessible edges.","Self-hosted Valhalla uses a versioned regional OSM graph; matching runs asynchronously."],
["Competition geometry","Represent what can be claimed and scored.","Runlete creates UUID-owned street edges, retains OSM lineage and scores accepted traversals in PostGIS."],
["Route analytics","Compare an activity with a route, segment or prior effort.","Uses matched geometry, coverage and confidence rather than raw proximity alone."]])
heading(doc,"OSM and Valhalla in simple terms",2)
ptext(doc,"OpenStreetMap (OSM) is the source map dataset describing roads, paths, trails and access attributes. Runlete imports regional, versioned OSM extracts and filters them to running-accessible geometry. Valhalla is the routing/matching engine that compares GPS samples against that graph. Mapbox displays the result; Runlete stores accepted facts and territory scores.")
for x in ["Activity distance at least 2.5 km.","Competitive quality checks pass and no unresolved anti-cheat flag exists.","At least 80% edge coverage and normalized matcher confidence of at least 0.85.","Unsupported-region, failed-match, processing and rejected states remain explicit; no raw trace is shown as verified territory."]: bullet(doc,x)

heading(doc,"7. Territory: not a circle, not latest trace wins")
ptext(doc,"Territory is road-edge control. A long straight run with turns can control a connected sequence of short verified edges. It is not a radius around the runner and not a freehand polygon.")
for x in ["Clean and map-match an eligible activity.","Translate matched geometry to Runlete street-edge UUIDs, splitting long ways into claimable edges, normally no longer than 100 m.","Keep the strongest qualifying traversal per athlete per local calendar day.","Score up to seven distinct days in the trailing 28 days, with decay: 100 × confidence × coverage × 0.5^(age_days / 14).","Assign personal control to the highest athlete score. Club control combines the five strongest currently attributed members.","Record claim, defence, loss, expiry and dispute history; recalculate when an activity is cropped, deleted, private, rejected or reprocessed."]: bullet(doc,x,True)
callout(doc,"Example","Two athletes run the same 4 km road. Athlete A has three high-confidence days; Athlete B has one excellent day. A can retain control through the rolling multi-day score even if B is faster once. Speed is a tie-breaker, not the whole system.","FFF8E1","7A5A00")

heading(doc,"8. Clubs, challenges and competition")
ptext(doc,"Runlete clubs are competitive groups, not social networks. Athletes may join multiple clubs but select exactly one primary competitive club. A future eligible activity is attributed to that club once; changing clubs never rewrites historical attribution.")
table(doc,["Capability","Behavior"],[
["Membership","Public clubs join immediately. Private clubs use request/approve/reject. Invitations expire; admins can remove, ban, archive, restore and transfer ownership."],
["Roles","Owner, admin and member. Each active club has exactly one owner. Mutations use shared policies and are audited."],
["Leaderboards","Weekly, monthly, season and all-time views for distance, moving time, consistency, territory, races and normalized per-active-member results."],
["Challenges","Distance, duration, run-count, consistency, territory-gain and fastest-effort challenges with provisional then verified results."],
["Races","Scheduled route, timezone, start window, capacity and eligibility. Results require route coverage and anti-cheat processing."],
["Club Activity","System-generated events for eligible runs, achievements, membership, challenge milestones, race results and territory changes. No posts, comments or chat."]])

heading(doc,"9. Routes, segments and heatmaps")
for x in ["Saved and generated routes with GPX import/export, route following and deviation state.","Curated and user-created running segments with geometry checks, automatic effort matching, PRs and leaderboards.","Personal privacy-filtered heatmaps and thresholded aggregate heatmaps.","Vector tiles for territory and map layers; mobile does not receive thousands of raw GeoJSON edges.","Visibility enforcement across public, club and private activities, exports, caches and tiles."]: bullet(doc,x)

heading(doc,"10. Workout generation: what the system actually does")
ptext(doc,"Runlete generates a program first and specific workouts second. This makes progression, spacing and sport relevance explicit instead of asking an LLM to invent a random daily session.")
callout(doc,"Core pipeline","Athlete input + sport demands + goal/phase + schedule/equipment/restrictions → four-week program structure → weekly structure → session template → exercise/modality selection → safety/workload validation → workout → completion/feedback → bounded adjustment","EAF4F4",TEAL)
heading(doc,"The five knowledge layers",2)
table(doc,["Layer","Stores","How it is used"],[
["1. Method library","Exercises, drills, conditioning, instructions, safety and media","Supplies actual selectable training methods."],
["2. Method effects","One primary effect, up to three secondary effects, role, costs and substitutions","Connects a method to the adaptation it can support."],
["3. Sport demands","Sport/event/role priorities and evidence-backed demand facts","Determines which qualities matter for the athlete’s context."],
["4. Prescription rules","Sets, reps, duration, intensity, recovery, progression, taper and stop rules","Controls valid doses and how they change."],
["5. Programs and templates","Program archetypes, phases, weeks and session recipes","Defines four-week progression and required session structure."]])
heading(doc,"Inputs collected from the athlete",2)
for x in ["Age confirmation (16+), primary sport, event/role and secondary sports.","Goal, target metric, target date, competition date and training phase.","Training age, recent volume, benchmarks and current readiness.","Available days, maximum duration, equipment and facility/environment.","External practices, matches/races, sleep/stress, pain or restrictions and feedback."]: bullet(doc,x)
heading(doc,"Where AI fits",2)
ptext(doc,"The current configuration uses OpenAI gpt-4o for workout selection and food-image analysis. The backend first retrieves only relevant template slots and approved method candidates. The model can choose from supplied IDs and bounded dose options; it cannot create an exercise, invent a dose, change sport priorities or bypass safety checks.")
for x in ["The backend chooses program length, phases, weekly structure, hard/easy spacing, candidate eligibility and prescription bounds.","The model receives a compact, relevant packet—not the entire research library or every exercise.","The response must contain valid supplied IDs, required slots, allowed units and in-range values. Invalid output fails closed and is not persisted.","The first two weeks can be materialized while the complete four-week structure remains stored.","Completion, RPE, readiness and pain feedback can modify future uncompleted work only within registered rules."]: bullet(doc,x)

heading(doc,"11. Nutrition and health around training")
ptext(doc,"Food-image analysis follows a confirmation model: the client uploads an image to Cloud Storage, an AI job returns candidate foods and nutrient ranges, deterministic validators reject malformed units or implausible values, and the athlete confirms or edits the meal. Confirmed meals—not raw AI guesses—drive summaries.")
ptext(doc,"Health data includes readiness, mood, sleep, stress, pain/restriction reports and assessments. It may conservatively reduce, replace, stop or refer a session. It does not diagnose injury or prescribe rehabilitation.")

heading(doc,"12. Admin Studio and operations")
ptext(doc,"Admin Studio is the operator’s control surface for the knowledge that powers generation. It should make normal editing safe without requiring UUIDs or raw JSON.")
for x in ["Exercises: view, add, edit, archive, attach media, choose controlled tags, edit effects and safety fields, link progressions through searchable selectors and preview generator eligibility.","Workout templates: view/import/create/edit four-week session structures, allowed method types, dosage bounds, substitutions and linked exercises.","Programs: view/import/create program archetypes, phases and weekly structures; preview schedules and linked templates.","Sports and goals: manage sport/event/role priorities and supporting rules that connect demands to relevant templates.","Users: view signup/subscription state and contact details with appropriate privacy and access controls.","Data imports: validate CSV/JSON payloads, show reconciliation/errors, commit atomically and make import history visible."]: bullet(doc,x)
callout(doc,"Release discipline","Editing content makes a new version. Publishing creates an immutable release. Existing plans remain pinned to their release; new generations use the newly published release.","FFF8E1","7A5A00")

heading(doc,"13. API and event boundaries")
ptext(doc,"Mobile and Admin Studio use generated clients from FastAPI OpenAPI. Public APIs are versioned under /api/v1 and use cursor pagination, stable problem responses, optimistic versions and idempotency keys for retriable mutations.")
table(doc,["Event","Producer","Consumers"],[
["activity.completed","Activity transaction + outbox","Analytics, map matching, segments, territory, club attribution, achievements and notifications."],
["training.session.completed","Training completion command","Readiness/load update, adaptation, streaks and feedback classification."],
["nutrition.analysis.requested","Signed-upload confirmation","Food AI worker, validation, confirmation UI and retention job."],
["content.release.published","Admin release transaction","Planner cache invalidation, release reports and future generation selection."],
["competition.result.verified","Competition processing","Leaderboards, challenge state, race result and Club Activity."]])

heading(doc,"14. Current status and important boundaries")
ptext(doc,"The repository contains a working foundation and an ongoing production rebuild. Treat these distinctions carefully when discussing the system:")
table(doc,["Area","Current direction","Still requires production hardening"],[
["Knowledge/admin","PostgreSQL-backed Admin Studio with exercise, template, program and sport-data imports.","Complete content releases, validation reports, production access and full catalogue coverage."],
["AI","OpenAI gpt-4o is configured for bounded workout selection and food analysis.","Real credentials, usage limits, prompt/version logging, outage handling and golden evaluations."],
["Running recorder","State-machine and activity domain are being rebuilt around SQL/GCS.","Real-device background GPS, resumable upload, workers and map-quality pilots."],
["Territory/clubs","Design is defined around OSM edges, Valhalla, rolling scores and SQL competition facts.","Regional graph deployment, anti-cheat, tile privacy and load testing."],
["Deployment","GCP target: Cloud Run, Cloud SQL/PostGIS, Pub/Sub, Storage, Redis and Valhalla compute.","Production IAM/secrets, Terraform apply, monitoring, backup restore and rollout gates."]])

heading(doc,"15. End-to-end example")
ptext(doc,"A recreational 10K runner trains four days per week, has football practice Tuesday and Thursday, owns dumbbells, and reports low readiness in week two.")
for x in ["Onboarding stores the 10K goal, four available days, football as secondary load, equipment and schedule.","Demand and priority data identify aerobic development, threshold work, running economy, tissue capacity and supporting strength; football practices are external load.","The planner builds a four-week structure and places hard running away from football congestion.","The selected template supplies blocks; approved method IDs and bounded doses are sent to gpt-4o for constrained selection.","Backend validation rejects anything that exceeds duration, duplicates a method improperly, violates spacing or uses unavailable equipment.","The runner reports low readiness. A registered rule reduces or replaces the next high-load session; it does not invent rehabilitation.","If the runner records GPS, the activity pipeline calculates facts, map-matches the route, and only then considers segment, territory, club and leaderboard contributions."]: bullet(doc,x,True)

heading(doc,"16. What a new contributor should remember")
for x in ["PostgreSQL is the authority for production domains; object storage holds large immutable files; Redis is not a database of record.","A raw GPS trace is evidence, not verified territory. Matching, quality and privacy processing happen first.","A club contribution is determined by primary-club attribution at activity time; historical attribution is immutable.","Sport demands choose priorities; templates choose structure; rules choose dose; methods supply the work; AI selects only within those boundaries.","Athlete-facing responses should be short and useful. Evidence, taxonomy scores, rule expressions and internal decision traces belong in Admin Studio and operations tooling.","Every asynchronous operation must be idempotent, observable, retryable and recoverable."]: bullet(doc,x)

doc.add_page_break()
heading(doc,"Appendix A. Glossary")
table(doc,["Term","Meaning"],[
["OSM","OpenStreetMap, the open geographic dataset imported into regional graph versions."],
["PostGIS","PostgreSQL extension for spatial geometry, indexes and geographic queries."],
["Valhalla","Self-hosted routing/map-matching engine that compares GPS samples to the OSM graph."],
["Map matching","Estimating which mapped road/path edges best explain a cleaned GPS stream."],
["Street edge","A Runlete-owned, UUID-identified claimable road/path segment with OSM lineage."],
["Primary club","The athlete’s one competitive club for future eligible activity attribution."],
["Content release","Immutable set of approved knowledge versions used by future plan generation."],
["Outbox","Transactional event record written with domain data so downstream processing is reliable."],
["Bounded AI","A model call whose candidate IDs, units, ranges and actions are supplied and validated by the backend."]])

for t in doc.tables:
    for row in t.rows:
        for c in row.cells:
            for p in c.paragraphs: p.paragraph_format.keep_together=True
            tcPr=c._tc.get_or_add_tcPr(); b=tcPr.first_child_found_in("w:tcBorders")
            if b is None: b=OxmlElement("w:tcBorders"); tcPr.append(b)
            for edge in ("top","left","bottom","right","insideH","insideV"):
                e=b.find(qn("w:"+edge))
                if e is None: e=OxmlElement("w:"+edge); b.append(e)
                e.set(qn("w:val"),"single"); e.set(qn("w:sz"),"3"); e.set(qn("w:color"),BORDER)
props=doc.core_properties; props.title="Runlete analysis"; props.subject="Runlete product and engineering reference"; props.author="Runlete"
doc.save(OUT); print(OUT)
